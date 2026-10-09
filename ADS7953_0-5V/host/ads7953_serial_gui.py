import queue
import threading
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

import serial
from openpyxl import Workbook, load_workbook
from serial.tools import list_ports


BAUD_RATE = 115200
TRIGGER_COMMAND = b"T"
FRAME_HEADER = b"\xA5\x5A"
PAYLOAD_LENGTH = 32
FULL_SCALE_MV = 5000
IV_TO_CH = (14, 12, 10, 8, 13, 15, 9, 11, 5, 7, 1, 3, 6, 4, 2, 0)


def append_excel_record(file_path, codes, remark):
    path = Path(file_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    new_sheet = False
    sheet_name = "采集记录_IV_0-5V"

    if path.exists():
        workbook = load_workbook(path)
        if sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
        else:
            sheet = workbook.create_sheet(sheet_name)
            new_sheet = True
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = sheet_name
        new_sheet = True

    if new_sheet:
        header = ["采集时间", "备注"]
        for iv in range(16):
            header.append(f"IV{iv}(mV)")
        sheet.append(header)
        sheet.freeze_panes = "C2"
        sheet.column_dimensions["A"].width = 21
        sheet.column_dimensions["B"].width = 32

    row = [datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3], remark]
    for code in codes:
        code &= 0x0FFF
        millivolts = (code * FULL_SCALE_MV + 2048) // 4096
        row.append(millivolts)
    sheet.append(row)
    workbook.save(path)
    workbook.close()


class SerialReader(threading.Thread):
    def __init__(self, port, output_queue, stop_event):
        super().__init__(daemon=True)
        self.port = port
        self.output_queue = output_queue
        self.stop_event = stop_event
        self.buffer = bytearray()

    def run(self):
        try:
            while not self.stop_event.is_set():
                chunk = self.port.read(self.port.in_waiting or 1)
                if chunk:
                    self.buffer.extend(chunk)
                    self.parse_frames()
        except serial.SerialException as exc:
            self.output_queue.put(("error", str(exc)))

    def parse_frames(self):
        while True:
            header_pos = self.buffer.find(FRAME_HEADER)
            if header_pos < 0:
                if self.buffer[-1:] == FRAME_HEADER[:1]:
                    self.buffer[:] = self.buffer[-1:]
                else:
                    self.buffer.clear()
                return

            if header_pos:
                del self.buffer[:header_pos]
            if len(self.buffer) < 3:
                return

            length = self.buffer[2]
            frame_length = 3 + length + 1
            if len(self.buffer) < frame_length:
                return

            frame = bytes(self.buffer[:frame_length])
            del self.buffer[:frame_length]
            checksum = 0
            for value in frame[2:-1]:
                checksum ^= value

            if checksum != frame[-1]:
                self.output_queue.put(("status", "收到校验错误的数据帧，已丢弃"))
                continue
            if length != PAYLOAD_LENGTH:
                self.output_queue.put(("status", f"收到未知长度的数据帧：{length}"))
                continue

            payload = frame[3:-1]
            codes = [payload[i] | (payload[i + 1] << 8)
                     for i in range(0, PAYLOAD_LENGTH, 2)]
            self.output_queue.put(("samples", codes))


class Ads7953App:
    def __init__(self, root):
        self.root = root
        self.root.title("ADS7953 16路电压采集（0-5V）")
        self.root.geometry("560x590")
        self.serial_port = None
        self.reader = None
        self.stop_event = None
        self.events = queue.Queue()
        self.awaiting_response = False
        self.pending_remark = ""

        toolbar = ttk.Frame(root, padding=10)
        toolbar.pack(fill="x")
        ttk.Label(toolbar, text="串口：").pack(side="left")
        self.port_box = ttk.Combobox(toolbar, width=18, state="readonly")
        self.port_box.pack(side="left", padx=(0, 6))
        ttk.Button(toolbar, text="刷新", command=self.refresh_ports).pack(side="left")
        self.connect_button = ttk.Button(toolbar, text="连接", command=self.toggle_connection)
        self.connect_button.pack(side="left", padx=6)
        self.trigger_button = ttk.Button(
            toolbar, text="触发一次采集", command=self.trigger_scan, state="disabled"
        )
        self.trigger_button.pack(side="right")

        record_frame = ttk.LabelFrame(root, text="本次采集记录", padding=8)
        record_frame.pack(fill="x", padx=10, pady=(0, 8))
        ttk.Label(record_frame, text="备注：").grid(row=0, column=0, sticky="nw")
        self.remark_text = tk.Text(record_frame, height=3, width=48)
        self.remark_text.grid(row=0, column=1, columnspan=2, sticky="ew")
        ttk.Label(record_frame, text="Excel：").grid(row=1, column=0, sticky="w", pady=(6, 0))
        default_excel = Path.home() / "Documents" / "ADS7953采集记录_0-5V.xlsx"
        self.excel_path = tk.StringVar(value=str(default_excel))
        ttk.Entry(record_frame, textvariable=self.excel_path).grid(
            row=1, column=1, sticky="ew", pady=(6, 0)
        )
        ttk.Button(record_frame, text="选择…", command=self.choose_excel).grid(
            row=1, column=2, padx=(6, 0), pady=(6, 0)
        )
        record_frame.columnconfigure(1, weight=1)

        columns = ("channel", "code", "mv", "voltage")
        self.table = ttk.Treeview(root, columns=columns, show="headings", height=16)
        self.table.heading("channel", text="通道")
        self.table.heading("code", text="原始码")
        self.table.heading("mv", text="毫伏")
        self.table.heading("voltage", text="电压(V)")
        self.table.column("channel", width=80, anchor="center")
        self.table.column("code", width=110, anchor="center")
        self.table.column("mv", width=110, anchor="center")
        self.table.column("voltage", width=130, anchor="center")
        self.table.pack(fill="both", expand=True, padx=10)
        self.rows = []
        for iv in range(16):
            item = self.table.insert("", "end", values=(f"IV{iv}", "---", "---", "---"))
            self.rows.append(item)

        self.status = tk.StringVar(value="未连接")
        ttk.Label(root, textvariable=self.status, padding=10).pack(fill="x")

        self.refresh_ports()
        self.root.after(50, self.process_events)
        self.root.protocol("WM_DELETE_WINDOW", self.close)

    def choose_excel(self):
        selected = filedialog.asksaveasfilename(
            title="选择采集记录文件",
            defaultextension=".xlsx",
            filetypes=(("Excel 工作簿", "*.xlsx"),),
            initialfile=Path(self.excel_path.get()).name,
        )
        if selected:
            self.excel_path.set(selected)

    def refresh_ports(self):
        ports = [port.device for port in list_ports.comports()]
        self.port_box["values"] = ports
        if ports and self.port_box.get() not in ports:
            self.port_box.current(0)
        if not ports:
            self.port_box.set("")

    def toggle_connection(self):
        if self.serial_port and self.serial_port.is_open:
            self.disconnect()
            return

        port_name = self.port_box.get()
        if not port_name:
            messagebox.showwarning("未选择串口", "请连接开发板并点击刷新。")
            return
        try:
            self.serial_port = serial.Serial(port_name, BAUD_RATE, timeout=0.1)
        except serial.SerialException as exc:
            messagebox.showerror("连接失败", str(exc))
            return

        self.stop_event = threading.Event()
        self.reader = SerialReader(self.serial_port, self.events, self.stop_event)
        self.reader.start()
        self.connect_button.config(text="断开")
        self.trigger_button.config(state="normal")
        self.status.set(f"已连接 {port_name}，{BAUD_RATE} baud，8N1")

    def disconnect(self):
        if self.stop_event:
            self.stop_event.set()
        if self.serial_port:
            try:
                self.serial_port.close()
            except serial.SerialException:
                pass
        self.serial_port = None
        self.awaiting_response = False
        self.pending_remark = ""
        self.connect_button.config(text="连接")
        self.trigger_button.config(state="disabled")
        self.status.set("未连接")

    def trigger_scan(self):
        try:
            self.pending_remark = self.remark_text.get("1.0", "end-1c")
            self.awaiting_response = True
            self.trigger_button.config(state="disabled")
            self.serial_port.reset_input_buffer()
            self.serial_port.write(TRIGGER_COMMAND)
            self.status.set("已发送触发命令，等待16路数据并写入Excel……")
        except (AttributeError, serial.SerialException) as exc:
            messagebox.showerror("发送失败", str(exc))
            self.disconnect()

    def process_events(self):
        while True:
            try:
                event_type, value = self.events.get_nowait()
            except queue.Empty:
                break

            if event_type == "samples":
                iv_codes = [value[channel] for channel in IV_TO_CH]
                for iv, code in enumerate(iv_codes):
                    code &= 0x0FFF
                    millivolts = (code * FULL_SCALE_MV + 2048) // 4096
                    self.table.item(
                        self.rows[iv],
                        values=(f"IV{iv}", f"0x{code:03X}", millivolts,
                                f"{millivolts / 1000:.3f}"),
                    )
                if self.awaiting_response:
                    try:
                        append_excel_record(self.excel_path.get(), iv_codes, self.pending_remark)
                    except (OSError, PermissionError, ValueError) as exc:
                        messagebox.showerror(
                            "Excel写入失败",
                            f"数据已收到，但无法写入Excel：\n{exc}\n\n"
                            "如果文件正被Excel打开，请先关闭文件后重新采集。",
                        )
                        self.status.set("数据已收到，但Excel写入失败")
                    else:
                        self.status.set(f"16路数据已写入：{self.excel_path.get()}")
                    self.awaiting_response = False
                    self.pending_remark = ""
                    if self.serial_port and self.serial_port.is_open:
                        self.trigger_button.config(state="normal")
            elif event_type == "status":
                self.status.set(value)
            elif event_type == "error":
                messagebox.showerror("串口错误", value)
                self.disconnect()

        self.root.after(50, self.process_events)

    def close(self):
        self.disconnect()
        self.root.destroy()


if __name__ == "__main__":
    window = tk.Tk()
    Ads7953App(window)
    window.mainloop()
