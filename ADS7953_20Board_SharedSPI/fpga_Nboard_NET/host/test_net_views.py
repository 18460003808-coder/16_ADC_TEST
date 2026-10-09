import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from ads7953_all_channels_NET_gui import AllChannelsApp, all_channel_rows
from ads7953_scope_NET_gui import ScopeApp, validate_scale_values


class NetViewTests(unittest.TestCase):
    def test_scope_scale_validation(self):
        self.assertEqual(
            (500, 50, -1.0, 6.0, 0.5),
            validate_scale_values("500", "50", "-1", "6", "0.5"))
        with self.assertRaises(ValueError):
            validate_scale_values("5", "1", "0", "5", "1")
        with self.assertRaises(ValueError):
            validate_scale_values("200", "20", "5", "0", "1")

    def test_global_iv_mapping(self):
        boards = {0: list(range(16)), 1: [100 + value for value in range(16)]}
        rows = list(all_channel_rows(boards, 2))
        self.assertEqual(32, len(rows))
        self.assertEqual((0, 0, 0, 14, 14), rows[0][:5])
        self.assertEqual((16, 1, 0, 14, 114), rows[16][:5])

    def test_all_channels_excel(self):
        boards = {0: list(range(16)), 1: [100 + value for value in range(16)]}
        with TemporaryDirectory() as directory:
            path = Path(directory) / "all.xlsx"
            app = AllChannelsApp.__new__(AllChannelsApp)
            app.export_path = path
            app._append_excel(2, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("IV0(mV)", sheet.cell(1, 4).value)
            self.assertEqual("IV16(mV)", sheet.cell(1, 20).value)
            self.assertEqual(35, sheet.max_column)

    def test_scope_excel_selected_channels(self):
        boards = {0: list(range(16)), 1: [100 + value for value in range(16)]}
        with TemporaryDirectory() as directory:
            path = Path(directory) / "scope.xlsx"
            app = ScopeApp.__new__(ScopeApp)
            app.export_path = path
            app._selected_channels = lambda: [0, 16]
            app._append_excel(3, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("IV0", sheet.cell(2, 3).value)
            self.assertEqual("IV16", sheet.cell(3, 3).value)


if __name__ == "__main__":
    unittest.main()
