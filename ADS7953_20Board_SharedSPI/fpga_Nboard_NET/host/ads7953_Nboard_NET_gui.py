"""ADS7953 N-board UDP acquisition tool."""

from __future__ import annotations

import queue
import socket
import threading
import time
import ipaddress
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook


FPGA_UDP_PORT = 5000
HOST_UDP_PORT = 5001
FPGA_IP = "192.168.1.50"
DEFAULT_BROADCAST_IP = "192.168.1.255"
MAX_BOARD_COUNT = 20
HEADER = b"\xA5\x5A"
VERSION = 1
PAYLOAD_LENGTH = 32
DATA_FRAME_LENGTH = 40
END_FRAME_LENGTH = 8
IV_TO_CHANNEL = (14, 12, 10, 8, 13, 15, 9, 11, 5, 7, 1, 3, 6, 4, 2, 0)


def local_ipv4_addresses() -> list[str]:
    """Return usable IPv4 addresses reported for this computer."""
    addresses: set[str] = set()
    try:
        infos = socket.getaddrinfo(
            socket.gethostname(), None, socket.AF_INET, socket.SOCK_DGRAM)
        addresses.update(info[4][0] for info in infos)
    except OSError:
        pass
    return sorted(
        address for address in addresses
        if not ipaddress.ip_address(address).is_loopback
        and not ipaddress.ip_address(address).is_link_local)


def discovery_candidates(addresses: list[str]) -> list[str]:
    """Keep addresses that can directly reach the fixed FPGA /24 subnet."""
    fpga_network = ipaddress.ip_network(f"{FPGA_IP}/24", strict=False)
    return [
        address for address in addresses
        if ipaddress.ip_address(address) in fpga_network
        and address != FPGA_IP
    ]


class FrameParser:
    def __init__(self) -> None:
        self.buffer = bytearray()

    def feed(self, data: bytes) -> list[tuple[str, int, int, list[int]]]:
        self.buffer.extend(data)
        frames: list[tuple[str, int, int, list[int]]] = []
        while True:
            start = self.buffer.find(HEADER)
            if start < 0:
                self.buffer[:] = self.buffer[-1:] if self.buffer[-1:] == HEADER[:1] else b""
                break
            if start:
                del self.buffer[:start]
            if len(self.buffer) < 5:
                break

            is_end = self.buffer[3] == 0xFF
            frame_length = END_FRAME_LENGTH if is_end else DATA_FRAME_LENGTH
            if len(self.buffer) < frame_length:
                break
            frame = bytes(self.buffer[:frame_length])

            if is_end:
                if (frame[2] != VERSION or not 1 <= frame[4] <= MAX_BOARD_COUNT or
                        self._xor(frame[2:7]) != frame[7]):
                    del self.buffer[0]
                    continue
                sequence = frame[5] | (frame[6] << 8)
                frames.append(("end", sequence, frame[4], []))
                del self.buffer[:frame_length]
                continue

            if (frame[2] != VERSION or frame[4] != PAYLOAD_LENGTH or
                    frame[3] >= MAX_BOARD_COUNT or self._xor(frame[2:39]) != frame[39]):
                del self.buffer[0]
                continue

            board_id = frame[3]
            sequence = frame[5] | (frame[6] << 8)
            codes = [
                frame[7 + 2 * channel] | ((frame[8 + 2 * channel] & 0x0F) << 8)
                for channel in range(16)
            ]
            frames.append(("data", sequence, board_id, codes))
            del self.buffer[:frame_length]
        return frames

    @staticmethod
    def _xor(data: bytes) -> int:
        result = 0
        for value in data:
            result ^= value
        return result


class Ads7953App:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("ADS7953 N板网口采集工具")
        self.root.geometry("900x710")
        self.udp_socket: socket.socket | None = None
        self.stop_event = threading.Event()
        self.continuous_stop_event = threading.Event()
        self.continuous_active = False
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.pending: dict[int, dict[int, list[int]]] = {}
        self.current_codes = {board: [0] * 16 for board in range(MAX_BOARD_COUNT)}
        self.current_sequence: int | None = None
        self.detected_board_count = 0
        self.discovery_active = False
        self.discovered_fpga_ip: str | None = None
        self.export_path = Path.cwd() / "ADS7953_NBoard_NET_data.xlsx"
        self._build_ui()
        self.root.after(50, self._poll_events)
        self.root.after(300, self._auto_discover)
        self.root.protocol("WM_DELETE_WINDOW", self._close)

    def _build_ui(self) -> None:
        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        self._build_network_controls(controls)
        self.trigger_button = ttk.Button(
            controls, text="采集", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="选择Excel", command=self._select_export).pack(side=tk.LEFT, padx=5)

        repeat_controls = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat_controls.pack(fill=tk.X)
        ttk.Label(repeat_controls, text="连续触发间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat_controls, textvariable=self.interval_var, width=8).pack(
            side=tk.LEFT, padx=(0, 8))
        ttk.Label(repeat_controls, text="采集时间(s，0=不限)：").pack(side=tk.LEFT)
        self.duration_var = tk.StringVar(value="10")
        ttk.Entry(repeat_controls, textvariable=self.duration_var, width=8).pack(
            side=tk.LEFT, padx=(0, 8))
        self.continuous_button = ttk.Button(
            repeat_controls, text="开始连续采集", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=5)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat_controls, text="连续采集时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)

        board_bar = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        board_bar.pack(fill=tk.X)
        ttk.Label(board_bar, text="查看板号：").pack(side=tk.LEFT)
        self.board_box = ttk.Combobox(
            board_bar, width=8, state="readonly",
            values=["0"])
        self.board_box.current(0)
        self.board_box.pack(side=tk.LEFT)
        self.board_box.bind("<<ComboboxSelected>>", lambda _event: self._refresh_table())
        self.sequence_text = tk.StringVar(value="采集序号：--")
        ttk.Label(board_bar, textvariable=self.sequence_text).pack(side=tk.LEFT, padx=20)

        self.status = tk.StringVar(value="未连接")
        ttk.Label(self.root, textvariable=self.status, padding=(8, 0, 8, 6)).pack(fill=tk.X)

        columns = ("iv", "channel", "code", "mv", "voltage")
        self.table = ttk.Treeview(self.root, columns=columns, show="headings", height=16)
        headings = ("接口", "ADC通道", "原始码", "毫伏", "电压(V)")
        widths = (100, 120, 120, 150, 150)
        for column, heading, width in zip(columns, headings, widths):
            self.table.heading(column, text=heading)
            self.table.column(column, width=width, anchor=tk.CENTER)
        self.table.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))
        for iv, channel in enumerate(IV_TO_CHANNEL):
            self.table.insert("", tk.END, iid=str(iv),
                              values=(f"IV{iv}", channel, 0, "0.00", "0.0000"))

    def _build_network_controls(self, parent: ttk.Frame) -> None:
        ttk.Label(parent, text="本机IP：").pack(side=tk.LEFT)
        self.bind_ip_var = tk.StringVar(value="0.0.0.0")
        ttk.Entry(parent, textvariable=self.bind_ip_var, width=14).pack(
            side=tk.LEFT, padx=(0, 5))
        ttk.Label(parent, text="目标广播IP：").pack(side=tk.LEFT)
        self.target_ip_var = tk.StringVar(value=DEFAULT_BROADCAST_IP)
        ttk.Entry(parent, textvariable=self.target_ip_var, width=16).pack(
            side=tk.LEFT, padx=(0, 5))
        self.connect_button = ttk.Button(
            parent, text="打开网口", command=self._toggle_connection)
        self.connect_button.pack(side=tk.LEFT, padx=5)
        self.discover_button = ttk.Button(
            parent, text="自动查找FPGA", command=self._auto_discover)
        self.discover_button.pack(side=tk.LEFT, padx=5)

    def _auto_discover(self) -> None:
        if self.discovery_active:
            return
        if self.udp_socket:
            return
        candidates = discovery_candidates(local_ipv4_addresses())
        if not candidates:
            self.status.set(
                "未找到192.168.1.x本机地址，请先把有线网卡配置为192.168.1.10/24")
            return
        self.discovery_active = True
        self.discover_button.configure(state=tk.DISABLED)
        self.connect_button.configure(state=tk.DISABLED)
        self.status.set("正在自动查找FPGA：" + "、".join(candidates))
        threading.Thread(
            target=self._discovery_worker, args=(candidates,), daemon=True).start()

    def _discovery_worker(self, candidates: list[str]) -> None:
        errors: list[str] = []
        for local_ip in candidates:
            udp_socket: socket.socket | None = None
            try:
                udp_socket = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
                udp_socket.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
                udp_socket.setsockopt(socket.SOL_SOCKET, socket.SO_BROADCAST, 1)
                udp_socket.bind((local_ip, HOST_UDP_PORT))
                udp_socket.settimeout(0.2)
                udp_socket.sendto(b"T", (DEFAULT_BROADCAST_IP, FPGA_UDP_PORT))
                parser = FrameParser()
                deadline = time.monotonic() + 1.5
                while time.monotonic() < deadline:
                    try:
                        data, sender = udp_socket.recvfrom(2048)
                    except socket.timeout:
                        continue
                    if sender[1] != FPGA_UDP_PORT:
                        continue
                    for frame_kind, _sequence, value, _codes in parser.feed(data):
                        if frame_kind == "end":
                            self.events.put((
                                "discovery_ok", (local_ip, sender[0], value)))
                            return
                errors.append(f"{local_ip}无响应")
            except OSError as exc:
                errors.append(f"{local_ip}: {exc}")
            finally:
                if udp_socket:
                    udp_socket.close()
        self.events.put(("discovery_failed", "；".join(errors)))

    def _toggle_connection(self) -> None:
        if self.udp_socket:
            self._disconnect()
            return
        bind_ip = self.bind_ip_var.get().strip()
        target_ip = self.target_ip_var.get().strip()
        try:
            socket.inet_aton(bind_ip)
            socket.inet_aton(target_ip)
        except OSError:
            messagebox.showwarning("参数错误", "本机IP或目标广播IP格式不正确。")
            return
        try:
            udp_socket = socket.socket(socket.AF_INET, socket.SOCK_DGRAM)
            udp_socket.setsockopt(socket.SOL_SOCKET, socket.SO_REUSEADDR, 1)
            udp_socket.setsockopt(socket.SOL_SOCKET, socket.SO_BROADCAST, 1)
            udp_socket.bind((bind_ip, HOST_UDP_PORT))
            udp_socket.settimeout(0.1)
            self.udp_socket = udp_socket
        except OSError as exc:
            messagebox.showerror("网口打开失败", str(exc))
            return
        self.stop_event.clear()
        threading.Thread(target=self._reader, daemon=True).start()
        self.connect_button.configure(text="关闭网口")
        self.discover_button.configure(state=tk.DISABLED)
        self.trigger_button.configure(state=tk.NORMAL)
        self.continuous_button.configure(state=tk.NORMAL)
        self.status.set(
            f"UDP已打开：本机端口{HOST_UDP_PORT}，目标{target_ip}:{FPGA_UDP_PORT}")

    def _disconnect(self) -> None:
        self._stop_continuous(update_status=False)
        self.stop_event.set()
        udp_socket = self.udp_socket
        self.udp_socket = None
        if udp_socket:
            try:
                udp_socket.close()
            except OSError:
                pass
        self.connect_button.configure(text="打开网口")
        self.discover_button.configure(state=tk.NORMAL)
        self.trigger_button.configure(state=tk.DISABLED)
        self.continuous_button.configure(state=tk.DISABLED)
        self.status.set("未连接")

    def _reader(self) -> None:
        parser = FrameParser()
        udp_socket = self.udp_socket
        try:
            while udp_socket and not self.stop_event.is_set():
                try:
                    data, _sender = udp_socket.recvfrom(2048)
                except socket.timeout:
                    continue
                for frame in parser.feed(data):
                    self.events.put(("frame", frame))
        except OSError as exc:
            if self.stop_event.is_set():
                return
            self.events.put(("error", str(exc)))

    def _trigger(self) -> None:
        if not self.udp_socket:
            return
        try:
            self.pending.clear()
            self.udp_socket.sendto(
                b"T", (self.target_ip_var.get().strip(), FPGA_UDP_PORT))
            self.trigger_button.configure(state=tk.DISABLED)
            self.status.set("正在采集，等待FPGA报告实际板卡数量……")
            self.root.after(6000, self._trigger_timeout)
        except OSError as exc:
            self._network_error(str(exc))

    def _toggle_continuous(self) -> None:
        if self.continuous_active:
            self._stop_continuous()
            return
        if not self.udp_socket:
            return
        try:
            interval_ms = int(self.interval_var.get())
            duration_s = float(self.duration_var.get())
        except ValueError:
            messagebox.showwarning("参数错误", "触发间隔和采集时间必须是数字。")
            return
        if not 20 <= interval_ms <= 60000:
            messagebox.showwarning("参数错误", "触发间隔范围为20～60000 ms。")
            return
        if not 0 <= duration_s <= 86400:
            messagebox.showwarning("参数错误", "采集时间范围为0～86400秒，0表示不限时。")
            return

        self.pending.clear()
        self.continuous_stop_event.clear()
        self.continuous_active = True
        self.trigger_button.configure(state=tk.DISABLED)
        self.continuous_button.configure(text="停止连续采集")
        duration_text = "不限时" if duration_s == 0 else f"持续 {duration_s:g} 秒"
        self.status.set(
            f"连续采集中：自动识别板数，每 {interval_ms} ms触发，{duration_text}")
        threading.Thread(
            target=self._continuous_writer,
            args=(interval_ms, duration_s), daemon=True).start()

    def _continuous_writer(self, interval_ms: int, duration_s: float) -> None:
        udp_socket = self.udp_socket
        target = (self.target_ip_var.get().strip(), FPGA_UDP_PORT)
        deadline = time.monotonic() + duration_s if duration_s > 0 else None
        try:
            while udp_socket and not self.continuous_stop_event.is_set():
                udp_socket.sendto(b"T", target)
                wait_seconds = interval_ms / 1000.0
                if deadline is not None:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0:
                        break
                    wait_seconds = min(wait_seconds, remaining)
                if self.continuous_stop_event.wait(wait_seconds):
                    return
            if not self.continuous_stop_event.is_set():
                self.events.put(("continuous_done", None))
        except OSError as exc:
            self.events.put(("error", str(exc)))

    def _stop_continuous(self, update_status: bool = True) -> None:
        self.continuous_stop_event.set()
        self.continuous_active = False
        if hasattr(self, "continuous_button"):
            self.continuous_button.configure(text="开始连续采集")
        if self.udp_socket:
            self.trigger_button.configure(state=tk.NORMAL)
            if update_status:
                self.status.set("连续采集已停止")

    def _trigger_timeout(self) -> None:
        if (not self.continuous_active and
                str(self.trigger_button["state"]) == tk.DISABLED and self.udp_socket):
            self.trigger_button.configure(state=tk.NORMAL)
            self.status.set("采集超时：请检查千兆网链路、IP配置、CS和共用SPI")

    def _poll_events(self) -> None:
        try:
            while True:
                kind, payload = self.events.get_nowait()
                if kind == "frame":
                    frame_kind, sequence, value, codes = payload  # type: ignore[misc]
                    if frame_kind == "data":
                        boards = self.pending.setdefault(sequence, {})
                        boards[value] = codes
                        self.status.set(f"序号 {sequence}：已收到 {len(boards)} 块板")
                    else:
                        board_count = value
                        boards = self.pending.get(sequence, {})
                        missing = [
                            board for board in range(board_count)
                            if board not in boards
                        ]
                        if missing:
                            self.status.set(
                                f"序号 {sequence}不完整，缺少板号："
                                + ",".join(str(board) for board in missing))
                            if not self.continuous_active:
                                self.trigger_button.configure(state=tk.NORMAL)
                            continue
                        self._complete_capture(sequence, boards, board_count)
                        self.pending = {
                            seq: values for seq, values in self.pending.items()
                            if seq > sequence
                        }
                elif kind == "continuous_done":
                    self._stop_continuous()
                    self.status.set("已达到设定采集时间，连续采集已停止")
                elif kind == "discovery_ok":
                    local_ip, fpga_ip, board_count = payload  # type: ignore[misc]
                    self.discovery_active = False
                    self.discovered_fpga_ip = fpga_ip
                    self.bind_ip_var.set(local_ip)
                    self.target_ip_var.set(DEFAULT_BROADCAST_IP)
                    self.connect_button.configure(state=tk.NORMAL)
                    self.discover_button.configure(state=tk.NORMAL)
                    self._toggle_connection()
                    if self.udp_socket:
                        self.status.set(
                            f"已发现FPGA {fpga_ip}：本机{local_ip}，"
                            f"检测到{board_count}块板，网口已自动打开")
                elif kind == "discovery_failed":
                    self.discovery_active = False
                    self.connect_button.configure(state=tk.NORMAL)
                    self.discover_button.configure(state=tk.NORMAL)
                    self.status.set(f"未发现FPGA：{payload}")
                else:
                    self._network_error(str(payload))
        except queue.Empty:
            pass
        self.root.after(50, self._poll_events)

    def _complete_capture(
            self, sequence: int, boards: dict[int, list[int]],
            board_count: int) -> None:
        self.detected_board_count = board_count
        for board in range(board_count):
            self.current_codes[board] = boards[board]
        self.board_box["values"] = [str(i) for i in range(board_count)]
        if int(self.board_box.get() or 0) >= board_count:
            self.board_box.current(0)
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self._refresh_table()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        mode = "连续采集" if self.continuous_active else "单次采集"
        self.status.set(
            f"序号 {sequence}：{mode}完成{board_count}板"
            f"共{board_count * 16}路{suffix}")

    def _refresh_table(self) -> None:
        board = int(self.board_box.get() or 0)
        codes = self.current_codes[board]
        for iv, channel in enumerate(IV_TO_CHANNEL):
            code = codes[channel]
            millivolts = code * 5000.0 / 4096.0
            self.table.item(str(iv), values=(
                f"IV{iv}", channel, code,
                f"{millivolts:.2f}", f"{millivolts / 1000.0:.4f}"))

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel输出：{self.export_path}")

    def _append_excel(
            self, sequence: int, boards: dict[int, list[int]],
            board_count: int) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "采集数据"
            sheet.append(["时间", "序号", "板号"] + [f"IV{i}(mV)" for i in range(16)])
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        for board in range(board_count):
            codes = boards[board]
            values = [codes[channel] * 5000.0 / 4096.0 for channel in IV_TO_CHANNEL]
            sheet.append([timestamp, sequence, board] + values)
        workbook.save(self.export_path)

    def _network_error(self, detail: str) -> None:
        self._disconnect()
        messagebox.showerror("网口错误", detail)

    def _close(self) -> None:
        self._disconnect()
        self.root.destroy()


def main() -> None:
    root = tk.Tk()
    Ads7953App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
