"""ADS7953 two-board shared-SPI validation tool."""

from __future__ import annotations

import queue
import threading
import time
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import serial
from serial.tools import list_ports
from openpyxl import Workbook


BAUD = 115200
HEADER = b"\xA5\x5A"
VERSION = 1
PAYLOAD_LENGTH = 32
FRAME_LENGTH = 40
BOARD_IDS = (0, 1)
IV_TO_CHANNEL = (14, 12, 10, 8, 13, 15, 9, 11, 5, 7, 1, 3, 6, 4, 2, 0)


class FrameParser:
    def __init__(self) -> None:
        self.buffer = bytearray()

    def feed(self, data: bytes) -> list[tuple[int, int, list[int]]]:
        self.buffer.extend(data)
        frames: list[tuple[int, int, list[int]]] = []
        while True:
            start = self.buffer.find(HEADER)
            if start < 0:
                if self.buffer[-1:] != HEADER[:1]:
                    self.buffer.clear()
                else:
                    self.buffer[:] = self.buffer[-1:]
                break
            if start:
                del self.buffer[:start]
            if len(self.buffer) < FRAME_LENGTH:
                break

            frame = bytes(self.buffer[:FRAME_LENGTH])
            if (frame[2] != VERSION or frame[4] != PAYLOAD_LENGTH or
                    frame[3] not in BOARD_IDS or
                    self._xor(frame[2:39]) != frame[39]):
                del self.buffer[0]
                continue

            board_id = frame[3]
            sequence = frame[5] | (frame[6] << 8)
            codes = [
                frame[7 + 2 * channel] |
                ((frame[8 + 2 * channel] & 0x0F) << 8)
                for channel in range(16)
            ]
            frames.append((sequence, board_id, codes))
            del self.buffer[:FRAME_LENGTH]
        return frames

    @staticmethod
    def _xor(data: bytes) -> int:
        result = 0
        for value in data:
            result ^= value
        return result


class Ads7953App:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("ADS7953 两板共用 SPI 验证工具")
        self.root.geometry("820x760")
        self.serial_port: serial.Serial | None = None
        self.stop_event = threading.Event()
        self.continuous_stop_event = threading.Event()
        self.continuous_active = False
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.pending: dict[int, dict[int, list[int]]] = {}
        self.current_codes = {board: [0] * 16 for board in BOARD_IDS}
        self.last_sequence: int | None = None
        self.export_path = Path.cwd() / "ADS7953_2Board_data.xlsx"
        self._build_ui()
        self._refresh_ports()
        self.root.after(50, self._poll_events)
        self.root.protocol("WM_DELETE_WINDOW", self._close)

    def _build_ui(self) -> None:
        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        ttk.Label(controls, text="串口：").pack(side=tk.LEFT)
        self.port_box = ttk.Combobox(controls, width=18, state="readonly")
        self.port_box.pack(side=tk.LEFT, padx=(0, 6))
        ttk.Button(controls, text="刷新", command=self._refresh_ports).pack(side=tk.LEFT)
        self.connect_button = ttk.Button(
            controls, text="连接", command=self._toggle_connection)
        self.connect_button.pack(side=tk.LEFT, padx=6)
        self.trigger_button = ttk.Button(
            controls, text="采集两板", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=6)

        repeat_controls = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat_controls.pack(fill=tk.X)
        ttk.Label(repeat_controls, text="连续触发间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat_controls, textvariable=self.interval_var, width=8).pack(
            side=tk.LEFT, padx=(0, 6))
        self.continuous_button = ttk.Button(
            repeat_controls, text="开始连续采集", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=6)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat_controls, text="连续采集时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)
        ttk.Button(controls, text="选择 Excel", command=self._select_export).pack(side=tk.LEFT)

        self.status = tk.StringVar(value="未连接")
        ttk.Label(self.root, textvariable=self.status, padding=(8, 0, 8, 6)).pack(fill=tk.X)

        columns = ("board", "iv", "channel", "code", "mv", "voltage")
        self.table = ttk.Treeview(self.root, columns=columns, show="headings", height=32)
        headings = ("板号", "接口", "ADC通道", "原始码", "毫伏", "电压(V)")
        widths = (65, 70, 90, 90, 100, 110)
        for column, heading, width in zip(columns, headings, widths):
            self.table.heading(column, text=heading)
            self.table.column(column, width=width, anchor=tk.CENTER)
        self.table.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))

        for board in BOARD_IDS:
            for iv, channel in enumerate(IV_TO_CHANNEL):
                self.table.insert("", tk.END, iid=f"{board}:{iv}",
                                  values=(board, f"IV{iv}", channel, 0, "0.00", "0.0000"))

    def _refresh_ports(self) -> None:
        ports = [port.device for port in list_ports.comports()]
        self.port_box["values"] = ports
        if ports and self.port_box.get() not in ports:
            self.port_box.current(0)

    def _toggle_connection(self) -> None:
        if self.serial_port and self.serial_port.is_open:
            self._disconnect()
            return
        port = self.port_box.get()
        if not port:
            messagebox.showwarning("提示", "没有可用串口，请连接设备后刷新。")
            return
        try:
            self.serial_port = serial.Serial(port, BAUD, timeout=0.1)
        except serial.SerialException as exc:
            messagebox.showerror("连接失败", str(exc))
            return
        self.stop_event.clear()
        threading.Thread(target=self._reader, daemon=True).start()
        self.connect_button.configure(text="断开")
        self.trigger_button.configure(state=tk.NORMAL)
        self.continuous_button.configure(state=tk.NORMAL)
        self.status.set(f"已连接 {port} @ {BAUD} baud")

    def _disconnect(self) -> None:
        self._stop_continuous(update_status=False)
        self.stop_event.set()
        port = self.serial_port
        self.serial_port = None
        if port:
            try:
                port.close()
            except serial.SerialException:
                pass
        self.connect_button.configure(text="连接")
        self.trigger_button.configure(state=tk.DISABLED)
        self.continuous_button.configure(state=tk.DISABLED)
        self.status.set("未连接")

    def _reader(self) -> None:
        parser = FrameParser()
        port = self.serial_port
        try:
            while port and port.is_open and not self.stop_event.is_set():
                data = port.read(port.in_waiting or 1)
                for frame in parser.feed(data):
                    self.events.put(("frame", frame))
        except serial.SerialException as exc:
            self.events.put(("error", str(exc)))

    def _trigger(self) -> None:
        if not self.serial_port or not self.serial_port.is_open:
            return
        try:
            self.serial_port.reset_input_buffer()
            self.serial_port.write(b"T")
            self.trigger_button.configure(state=tk.DISABLED)
            self.status.set("正在依次采集板 0、板 1……")
            self.root.after(3000, self._trigger_timeout)
        except serial.SerialException as exc:
            self._serial_error(str(exc))

    def _toggle_continuous(self) -> None:
        if self.continuous_active:
            self._stop_continuous()
            return
        if not self.serial_port or not self.serial_port.is_open:
            return
        try:
            interval_ms = int(self.interval_var.get())
        except ValueError:
            messagebox.showwarning("参数错误", "触发间隔必须是整数毫秒。")
            return
        if not 20 <= interval_ms <= 60000:
            messagebox.showwarning("参数错误", "触发间隔范围为20～60000 ms。")
            return

        try:
            self.serial_port.reset_input_buffer()
        except serial.SerialException as exc:
            self._serial_error(str(exc))
            return
        self.pending.clear()
        self.continuous_stop_event.clear()
        self.continuous_active = True
        self.trigger_button.configure(state=tk.DISABLED)
        self.continuous_button.configure(text="停止连续采集")
        self.status.set(f"连续采集中：每 {interval_ms} ms 发送一次触发")
        threading.Thread(
            target=self._continuous_writer, args=(interval_ms,), daemon=True).start()

    def _continuous_writer(self, interval_ms: int) -> None:
        port = self.serial_port
        try:
            while port and port.is_open and not self.continuous_stop_event.is_set():
                port.write(b"T")
                if self.continuous_stop_event.wait(interval_ms / 1000.0):
                    break
        except serial.SerialException as exc:
            self.events.put(("error", str(exc)))

    def _stop_continuous(self, update_status: bool = True) -> None:
        self.continuous_stop_event.set()
        self.continuous_active = False
        if hasattr(self, "continuous_button"):
            self.continuous_button.configure(text="开始连续采集")
        if self.serial_port and self.serial_port.is_open:
            self.trigger_button.configure(state=tk.NORMAL)
            if update_status:
                self.status.set("连续采集已停止")

    def _trigger_timeout(self) -> None:
        if (not self.continuous_active and
                str(self.trigger_button["state"]) == tk.DISABLED and self.serial_port):
            self.trigger_button.configure(state=tk.NORMAL)
            self.status.set("采集超时：请检查 CS、共用 SPI 和串口连接")

    def _poll_events(self) -> None:
        try:
            while True:
                kind, payload = self.events.get_nowait()
                if kind == "frame":
                    sequence, board, codes = payload  # type: ignore[misc]
                    boards = self.pending.setdefault(sequence, {})
                    boards[board] = codes
                    if all(board_id in boards for board_id in BOARD_IDS):
                        self._display_capture(sequence, boards)
                        self.pending = {
                            seq: values for seq, values in self.pending.items()
                            if seq > sequence
                        }
                else:
                    self._serial_error(str(payload))
        except queue.Empty:
            pass
        self.root.after(50, self._poll_events)

    def _display_capture(self, sequence: int, boards: dict[int, list[int]]) -> None:
        for board in BOARD_IDS:
            self.current_codes[board] = boards[board]
            for iv, channel in enumerate(IV_TO_CHANNEL):
                code = boards[board][channel]
                millivolts = code * 5000.0 / 4095.0
                self.table.item(f"{board}:{iv}", values=(
                    board, f"IV{iv}", channel, code,
                    f"{millivolts:.2f}", f"{millivolts / 1000.0:.4f}"))
        self.last_sequence = sequence
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        mode = "连续采集" if self.continuous_active else "单次采集"
        self.status.set(f"序号 {sequence}：{mode}完成两板32路{suffix}")

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel 输出：{self.export_path}")

    def _append_excel(self, sequence: int, boards: dict[int, list[int]]) -> None:
        if self.export_path.exists():
            from openpyxl import load_workbook
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "采集数据"
            sheet.append(["时间", "序号", "板号"] + [f"IV{i}(mV)" for i in range(16)])
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        for board in BOARD_IDS:
            codes = boards[board]
            values = [codes[channel] * 5000.0 / 4095.0 for channel in IV_TO_CHANNEL]
            sheet.append([timestamp, sequence, board] + values)
        workbook.save(self.export_path)

    def _serial_error(self, detail: str) -> None:
        self._disconnect()
        messagebox.showerror("串口错误", detail)

    def _close(self) -> None:
        self._disconnect()
        self.root.destroy()


def main() -> None:
    root = tk.Tk()
    Ads7953App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
