import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from ads7953_all_channels_gui import AllChannelsApp
from ads7953_all_channels_gui import all_channel_rows
from ads7953_scope_gui import ScopeApp


class AllChannelMappingTests(unittest.TestCase):
    def test_global_iv_mapping_across_boards(self):
        boards = {
            0: list(range(16)),
            1: [100 + value for value in range(16)],
        }
        rows = list(all_channel_rows(boards, 2))
        self.assertEqual(32, len(rows))
        self.assertEqual((0, 0, 0, 14, 14), rows[0][:5])
        self.assertEqual((15, 0, 15, 0, 0), rows[15][:5])
        self.assertEqual((16, 1, 0, 14, 114), rows[16][:5])
        self.assertEqual((31, 1, 15, 0, 100), rows[31][:5])

    def test_all_channels_excel_uses_one_row_and_global_iv_headers(self):
        boards = {0: list(range(16)), 1: [100 + value for value in range(16)]}
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "all.xlsx"
            app = AllChannelsApp.__new__(AllChannelsApp)
            app.export_path = path
            app._append_excel(7, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("IV0(mV)", sheet.cell(1, 4).value)
            self.assertEqual("IV16(mV)", sheet.cell(1, 20).value)
            self.assertEqual(35, sheet.max_column)
            self.assertEqual(2, sheet.max_row)

    def test_scope_excel_exports_only_selected_global_channels(self):
        boards = {0: list(range(16)), 1: [100 + value for value in range(16)]}
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "scope.xlsx"
            app = ScopeApp.__new__(ScopeApp)
            app.export_path = path
            app._selected_channels = lambda: [0, 16]
            app._append_excel(8, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("IV0", sheet.cell(2, 3).value)
            self.assertEqual("IV16", sheet.cell(3, 3).value)
            self.assertEqual(3, sheet.max_row)


if __name__ == "__main__":
    unittest.main()
