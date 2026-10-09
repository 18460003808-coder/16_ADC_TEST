"""ADS7953 N-board acquisition tool with all channels on one page."""

from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, ttk

from openpyxl import Workbook, load_workbook

from ads7953_Nboard_gui import Ads7953App, BAUD, IV_TO_CHANNEL


def all_channel_rows(boards: dict[int, list[int]], board_count: int):
    """Return rows in global IV order: board 0 IV0..15, board 1 IV16..31..."""
    for board in range(board_count):
        codes = boards[board]
        for local_iv, adc_channel in enumerate(IV_TO_CHANNEL):
            global_iv = board * 16 + local_iv
            code = codes[adc_channel]
            millivolts = code * 5000.0 / 4096.0
            yield global_iv, board, local_iv, adc_channel, code, millivolts


class AllChannelsApp(Ads7953App):
    def __init__(self, root: tk.Tk) -> None:
        super().__init__(root)
        self.export_path = Path.cwd() / "ADS7953_AllChannels_data.xlsx"

    def _build_ui(self) -> None:
        self.root.title("ADS7953 N板全通道采集工具")
        self.root.geometry("980x760")

        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        ttk.Label(controls, text="串口：").pack(side=tk.LEFT)
        self.port_box = ttk.Combobox(controls, width=16, state="readonly")
        self.port_box.pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(controls, text="刷新", command=self._refresh_ports).pack(side=tk.LEFT)
        self.connect_button = ttk.Button(controls, text="连接", command=self._toggle_connection)
        self.connect_button.pack(side=tk.LEFT, padx=5)
        self.trigger_button = ttk.Button(
            controls, text="采集", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="选择Excel", command=self._select_export).pack(side=tk.LEFT, padx=5)

        repeat = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat.pack(fill=tk.X)
        ttk.Label(repeat, text="连续触发间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat, textvariable=self.interval_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(repeat, text="采集时间(s，0=不限)：").pack(side=tk.LEFT)
        self.duration_var = tk.StringVar(value="10")
        ttk.Entry(repeat, textvariable=self.duration_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        self.continuous_button = ttk.Button(
            repeat, text="开始连续采集", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=5)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat, text="连续采集时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)

        info = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        info.pack(fill=tk.X)
        self.sequence_text = tk.StringVar(value="采集序号：--")
        self.channel_count_text = tk.StringVar(value="等待识别板卡数量")
        ttk.Label(info, textvariable=self.sequence_text).pack(side=tk.LEFT)
        ttk.Label(info, textvariable=self.channel_count_text).pack(side=tk.LEFT, padx=20)

        self.status = tk.StringVar(value="未连接")
        ttk.Label(self.root, textvariable=self.status, padding=(8, 0, 8, 6)).pack(fill=tk.X)

        table_frame = ttk.Frame(self.root, padding=(8, 0, 8, 8))
        table_frame.pack(fill=tk.BOTH, expand=True)
        columns = ("iv", "board", "local_iv", "channel", "code", "mv", "voltage")
        self.table = ttk.Treeview(table_frame, columns=columns, show="headings")
        headings = ("全局接口", "板号", "板内接口", "ADC通道", "原始码", "毫伏", "电压(V)")
        widths = (100, 70, 100, 100, 100, 130, 130)
        for column, heading, width in zip(columns, headings, widths):
            self.table.heading(column, text=heading)
            self.table.column(column, width=width, anchor=tk.CENTER)
        scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.table.yview)
        self.table.configure(yscrollcommand=scrollbar.set)
        self.table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)

    def _complete_capture(
            self, sequence: int, boards: dict[int, list[int]],
            board_count: int) -> None:
        self.detected_board_count = board_count
        for board in range(board_count):
            self.current_codes[board] = boards[board]
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self.channel_count_text.set(f"{board_count}块板，共{board_count * 16}路")
        self._refresh_table()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        mode = "连续采集" if self.continuous_active else "单次采集"
        self.status.set(f"序号 {sequence}：{mode}完成，{board_count * 16}路{suffix}")

    def _refresh_table(self) -> None:
        self.table.delete(*self.table.get_children())
        boards = {
            board: self.current_codes[board]
            for board in range(self.detected_board_count)
        }
        for global_iv, board, local_iv, adc_channel, code, millivolts in all_channel_rows(
                boards, self.detected_board_count):
            self.table.insert("", tk.END, values=(
                f"IV{global_iv}", board, f"IV{local_iv}", adc_channel, code,
                f"{millivolts:.2f}", f"{millivolts / 1000.0:.4f}"))

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择全通道数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel输出：{self.export_path}")

    def _append_excel(
            self, sequence: int, boards: dict[int, list[int]],
            board_count: int) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "全通道采集数据"
            sheet.append(
                ["时间", "序号", "板卡数量"]
                + [f"IV{iv}(mV)" for iv in range(board_count * 16)])
        values = [row[5] for row in all_channel_rows(boards, board_count)]
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        sheet.append([timestamp, sequence, board_count] + values)
        workbook.save(self.export_path)


def main() -> None:
    root = tk.Tk()
    AllChannelsApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
