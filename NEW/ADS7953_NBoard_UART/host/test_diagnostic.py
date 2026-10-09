import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

from ads7953_Nboard_gui import FrameParser
from ads7953_diagnostic_gui import DiagnosticApp, DiagnosticFrameParser


def xor(data: bytes) -> int:
    value = 0
    for byte in data:
        value ^= byte
    return value


def data_frame(sequence: int, board: int, spi_error: bool) -> bytes:
    body = bytearray([1, board, 32, sequence & 0xFF, sequence >> 8])
    for channel in range(16):
        code = board * 0x100 + channel
        high = (code >> 8) & 0x0F
        if channel == 0 and spi_error:
            high |= 0x80
        body.extend((code & 0xFF, high))
    return b"\xA5\x5A" + body + bytes([xor(body)])


class DiagnosticFrameParserTest(unittest.TestCase):
    def test_normal_frame_is_not_flagged(self) -> None:
        frame = DiagnosticFrameParser().feed(data_frame(7, 2, False))[0]
        self.assertEqual(("data", 7, 2), frame[:3])
        self.assertFalse(frame[4])

    def test_error_flag_is_reported_without_changing_adc_code(self) -> None:
        frame = DiagnosticFrameParser().feed(data_frame(8, 3, True))[0]
        self.assertTrue(frame[4])
        self.assertEqual(0x300, frame[3][0])

    def test_existing_parser_accepts_flagged_frame(self) -> None:
        frame = FrameParser().feed(data_frame(8, 3, True))[0]
        self.assertEqual(("data", 8, 3), frame[:3])
        self.assertEqual(0x300, frame[3][0])

    def test_excel_records_and_highlights_error(self) -> None:
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "diagnostic.xlsx"
            app = DiagnosticApp.__new__(DiagnosticApp)
            app.export_path = path
            app._append_excel(9, {0: ([0] * 16, True)}, 1)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("读取状态", sheet.cell(1, 4).value)
            self.assertEqual("读取异常（32字超时）", sheet.cell(2, 4).value)
            self.assertEqual("00C00000", sheet.cell(2, 4).font.color.rgb)


if __name__ == "__main__":
    unittest.main()
