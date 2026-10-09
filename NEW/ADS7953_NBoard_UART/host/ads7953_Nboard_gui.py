"""ADS7953 N-board shared-SPI acquisition tool."""

from __future__ import annotations

import queue
import threading
import time
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

import serial
from serial.tools import list_ports
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill


BAUD = 115200
MAX_BOARD_COUNT = 20
HEADER = b"\xA5\x5A"
VERSION = 1
PAYLOAD_LENGTH = 32
DATA_FRAME_LENGTH = 40
END_FRAME_LENGTH = 8
IV_TO_CHANNEL = (14, 12, 10, 8, 13, 15, 9, 11, 5, 7, 1, 3, 6, 4, 2, 0)


class FrameParser:
    def __init__(self) -> None:
        self.buffer = bytearray()

    def feed(self, data: bytes) -> list[tuple[str, int, int, list[int], bool]]:
        self.buffer.extend(data)
        frames: list[tuple[str, int, int, list[int], bool]] = []
        while True:
            start = self.buffer.find(HEADER)
            if start < 0:
                self.buffer[:] = self.buffer[-1:] if self.buffer[-1:] == HEADER[:1] else b""
                break
            if start:
                del self.buffer[:start]
            if len(self.buffer) < 5:
                break

            is_end = self.buffer[3] == 0xFF
            frame_length = END_FRAME_LENGTH if is_end else DATA_FRAME_LENGTH
            if len(self.buffer) < frame_length:
                break
            frame = bytes(self.buffer[:frame_length])

            if is_end:
                if (frame[2] != VERSION or not 1 <= frame[4] <= MAX_BOARD_COUNT or
                        self._xor(frame[2:7]) != frame[7]):
                    del self.buffer[0]
                    continue
                sequence = frame[5] | (frame[6] << 8)
                frames.append(("end", sequence, frame[4], [], False))
                del self.buffer[:frame_length]
                continue

            if (frame[2] != VERSION or frame[4] != PAYLOAD_LENGTH or
                    frame[3] >= MAX_BOARD_COUNT or self._xor(frame[2:39]) != frame[39]):
                del self.buffer[0]
                continue

            board_id = frame[3]
            sequence = frame[5] | (frame[6] << 8)
            codes = [
                frame[7 + 2 * channel] | ((frame[8 + 2 * channel] & 0x0F) << 8)
                for channel in range(16)
            ]
            spi_error = bool(frame[8] & 0x80)
            frames.append(("data", sequence, board_id, codes, spi_error))
            del self.buffer[:frame_length]
        return frames

    @staticmethod
    def _xor(data: bytes) -> int:
        result = 0
        for value in data:
            result ^= value
        return result


class Ads7953App:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("ADS7953 N板自动识别采集工具")
        self.root.geometry("900x710")
        self.serial_port: serial.Serial | None = None
        self.stop_event = threading.Event()
        self.continuous_stop_event = threading.Event()
        self.continuous_active = False
        self.events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.pending: dict[int, dict[int, tuple[list[int], bool]]] = {}
        self.current_codes = {board: [0] * 16 for board in range(MAX_BOARD_COUNT)}
        self.current_errors = {board: False for board in range(MAX_BOARD_COUNT)}
        self.current_sequence: int | None = None
        self.detected_board_count = 0
        self.export_path = Path.cwd() / "ADS7953_NBoard_data.xlsx"
        self._build_ui()
        self._refresh_ports()
        self.root.after(50, self._poll_events)
        self.root.protocol("WM_DELETE_WINDOW", self._close)

    def _build_ui(self) -> None:
        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        ttk.Label(controls, text="串口：").pack(side=tk.LEFT)
        self.port_box = ttk.Combobox(controls, width=16, state="readonly")
        self.port_box.pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(controls, text="刷新", command=self._refresh_ports).pack(side=tk.LEFT)
        self.connect_button = ttk.Button(controls, text="连接", command=self._toggle_connection)
        self.connect_button.pack(side=tk.LEFT, padx=5)
        self.trigger_button = ttk.Button(
            controls, text="采集", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="选择Excel", command=self._select_export).pack(side=tk.LEFT, padx=5)

        repeat_controls = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat_controls.pack(fill=tk.X)
        ttk.Label(repeat_controls, text="连续触发间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat_controls, textvariable=self.interval_var, width=8).pack(
            side=tk.LEFT, padx=(0, 8))
        ttk.Label(repeat_controls, text="采集时间(s，0=不限)：").pack(side=tk.LEFT)
        self.duration_var = tk.StringVar(value="10")
        ttk.Entry(repeat_controls, textvariable=self.duration_var, width=8).pack(
            side=tk.LEFT, padx=(0, 8))
        self.continuous_button = ttk.Button(
            repeat_controls, text="开始连续采集", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=5)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat_controls, text="连续采集时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)

        board_bar = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        board_bar.pack(fill=tk.X)
        ttk.Label(board_bar, text="查看板号：").pack(side=tk.LEFT)
        self.board_box = ttk.Combobox(
            board_bar, width=8, state="readonly",
            values=["0"])
        self.board_box.current(0)
        self.board_box.pack(side=tk.LEFT)
        self.board_box.bind("<<ComboboxSelected>>", lambda _event: self._refresh_table())
        self.sequence_text = tk.StringVar(value="采集序号：--")
        ttk.Label(board_bar, textvariable=self.sequence_text).pack(side=tk.LEFT, padx=20)

        self.status = tk.StringVar(value="未连接")
        ttk.Label(self.root, textvariable=self.status, padding=(8, 0, 8, 6)).pack(fill=tk.X)

        columns = ("iv", "channel", "code", "mv", "voltage")
        self.table = ttk.Treeview(self.root, columns=columns, show="headings", height=16)
        headings = ("接口", "ADC通道", "原始码", "毫伏", "电压(V)")
        widths = (100, 120, 120, 150, 150)
        for column, heading, width in zip(columns, headings, widths):
            self.table.heading(column, text=heading)
            self.table.column(column, width=width, anchor=tk.CENTER)
        self.table.pack(fill=tk.BOTH, expand=True, padx=8, pady=(0, 8))
        for iv, channel in enumerate(IV_TO_CHANNEL):
            self.table.insert("", tk.END, iid=str(iv),
                              values=(f"IV{iv}", channel, 0, "0.00", "0.0000"))
        self.table.tag_configure(
            "spi_error", foreground="#C00000", background="#FFE6E6")

    def _refresh_ports(self) -> None:
        ports = [port.device for port in list_ports.comports()]
        self.port_box["values"] = ports
        if ports and self.port_box.get() not in ports:
            self.port_box.current(0)

    def _toggle_connection(self) -> None:
        if self.serial_port and self.serial_port.is_open:
            self._disconnect()
            return
        port = self.port_box.get()
        if not port:
            messagebox.showwarning("提示", "没有可用串口，请连接设备后刷新。")
            return
        try:
            self.serial_port = serial.Serial(port, BAUD, timeout=0.1)
        except serial.SerialException as exc:
            messagebox.showerror("连接失败", str(exc))
            return
        self.stop_event.clear()
        threading.Thread(target=self._reader, daemon=True).start()
        self.connect_button.configure(text="断开")
        self.trigger_button.configure(state=tk.NORMAL)
        self.continuous_button.configure(state=tk.NORMAL)
        self.status.set(f"已连接 {port} @ {BAUD} baud")

    def _disconnect(self) -> None:
        self._stop_continuous(update_status=False)
        self.stop_event.set()
        port = self.serial_port
        self.serial_port = None
        if port:
            try:
                port.close()
            except serial.SerialException:
                pass
        self.connect_button.configure(text="连接")
        self.trigger_button.configure(state=tk.DISABLED)
        self.continuous_button.configure(state=tk.DISABLED)
        self.status.set("未连接")

    def _reader(self) -> None:
        parser = FrameParser()
        port = self.serial_port
        try:
            while port and port.is_open and not self.stop_event.is_set():
                data = port.read(port.in_waiting or 1)
                for frame in parser.feed(data):
                    self.events.put(("frame", frame))
        except serial.SerialException as exc:
            self.events.put(("error", str(exc)))

    def _trigger(self) -> None:
        if not self.serial_port or not self.serial_port.is_open:
            return
        try:
            self.serial_port.reset_input_buffer()
            self.pending.clear()
            self.serial_port.write(b"T")
            self.trigger_button.configure(state=tk.DISABLED)
            self.status.set("正在采集，等待FPGA报告实际板卡数量……")
            self.root.after(6000, self._trigger_timeout)
        except serial.SerialException as exc:
            self._serial_error(str(exc))

    def _toggle_continuous(self) -> None:
        if self.continuous_active:
            self._stop_continuous()
            return
        if not self.serial_port or not self.serial_port.is_open:
            return
        try:
            interval_ms = int(self.interval_var.get())
            duration_s = float(self.duration_var.get())
        except ValueError:
            messagebox.showwarning("参数错误", "触发间隔和采集时间必须是数字。")
            return
        if not 20 <= interval_ms <= 60000:
            messagebox.showwarning("参数错误", "触发间隔范围为20～60000 ms。")
            return
        if not 0 <= duration_s <= 86400:
            messagebox.showwarning("参数错误", "采集时间范围为0～86400秒，0表示不限时。")
            return

        try:
            self.serial_port.reset_input_buffer()
        except serial.SerialException as exc:
            self._serial_error(str(exc))
            return
        self.pending.clear()
        self.continuous_stop_event.clear()
        self.continuous_active = True
        self.trigger_button.configure(state=tk.DISABLED)
        self.continuous_button.configure(text="停止连续采集")
        duration_text = "不限时" if duration_s == 0 else f"持续 {duration_s:g} 秒"
        self.status.set(
            f"连续采集中：自动识别板数，每 {interval_ms} ms触发，{duration_text}")
        threading.Thread(
            target=self._continuous_writer,
            args=(interval_ms, duration_s), daemon=True).start()

    def _continuous_writer(self, interval_ms: int, duration_s: float) -> None:
        port = self.serial_port
        deadline = time.monotonic() + duration_s if duration_s > 0 else None
        try:
            while port and port.is_open and not self.continuous_stop_event.is_set():
                port.write(b"T")
                wait_seconds = interval_ms / 1000.0
                if deadline is not None:
                    remaining = deadline - time.monotonic()
                    if remaining <= 0:
                        break
                    wait_seconds = min(wait_seconds, remaining)
                if self.continuous_stop_event.wait(wait_seconds):
                    return
            if not self.continuous_stop_event.is_set():
                self.events.put(("continuous_done", None))
        except serial.SerialException as exc:
            self.events.put(("error", str(exc)))

    def _stop_continuous(self, update_status: bool = True) -> None:
        self.continuous_stop_event.set()
        self.continuous_active = False
        if hasattr(self, "continuous_button"):
            self.continuous_button.configure(text="开始连续采集")
        if self.serial_port and self.serial_port.is_open:
            self.trigger_button.configure(state=tk.NORMAL)
            if update_status:
                self.status.set("连续采集已停止")

    def _trigger_timeout(self) -> None:
        if (not self.continuous_active and
                str(self.trigger_button["state"]) == tk.DISABLED and self.serial_port):
            self.trigger_button.configure(state=tk.NORMAL)
            self.status.set("采集超时：请检查缺失板卡的CS、共用SPI和串口连接")

    def _poll_events(self) -> None:
        try:
            while True:
                kind, payload = self.events.get_nowait()
                if kind == "frame":
                    frame_kind, sequence, value, codes, spi_error = payload  # type: ignore[misc]
                    if frame_kind == "data":
                        boards = self.pending.setdefault(sequence, {})
                        boards[value] = (codes, spi_error)
                        self.status.set(f"序号 {sequence}：已收到 {len(boards)} 块板")
                    else:
                        board_count = value
                        boards = self.pending.get(sequence, {})
                        missing = [
                            board for board in range(board_count)
                            if board not in boards
                        ]
                        if missing:
                            self.status.set(
                                f"序号 {sequence}不完整，缺少板号："
                                + ",".join(str(board) for board in missing))
                            if not self.continuous_active:
                                self.trigger_button.configure(state=tk.NORMAL)
                            continue
                        self._complete_capture(sequence, boards, board_count)
                        self.pending = {
                            seq: values for seq, values in self.pending.items()
                            if seq > sequence
                        }
                elif kind == "continuous_done":
                    self._stop_continuous()
                    self.status.set("已达到设定采集时间，连续采集已停止")
                else:
                    self._serial_error(str(payload))
        except queue.Empty:
            pass
        self.root.after(50, self._poll_events)

    def _complete_capture(
            self, sequence: int, boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        self.detected_board_count = board_count
        for board in range(board_count):
            self.current_codes[board], self.current_errors[board] = boards[board]
        self.board_box["values"] = [str(i) for i in range(board_count)]
        if int(self.board_box.get() or 0) >= board_count:
            self.board_box.current(0)
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self._refresh_table()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        failed = [str(board) for board in range(board_count) if self.current_errors[board]]
        if failed:
            self.status.set(
                f"序号 {sequence}：板号 {','.join(failed)} SPI读取异常（32字超时）{suffix}")
        else:
            mode = "连续采集" if self.continuous_active else "单次采集"
            self.status.set(
                f"序号 {sequence}：{mode}完成{board_count}板"
                f"共{board_count * 16}路，全部正常{suffix}")

    def _refresh_table(self) -> None:
        board = int(self.board_box.get() or 0)
        codes = self.current_codes[board]
        row_tags = ("spi_error",) if self.current_errors[board] else ()
        for iv, channel in enumerate(IV_TO_CHANNEL):
            code = codes[channel]
            millivolts = code * 5000.0 / 4096.0
            self.table.item(
                str(iv), values=(f"IV{iv}", channel, code,
                                 f"{millivolts:.2f}", f"{millivolts / 1000.0:.4f}"),
                tags=row_tags)

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel输出：{self.export_path}")

    def _append_excel(
            self, sequence: int, boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "采集数据"
            sheet.append(
                ["时间", "序号", "板号", "读取状态"]
                + [f"IV{i}(mV)" for i in range(16)])
        if sheet.cell(1, 4).value != "读取状态":
            sheet.insert_cols(4)
            sheet.cell(1, 4, "读取状态")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        error_fill = PatternFill(fill_type="solid", fgColor="FFE6E6")
        error_font = Font(color="C00000")
        for board in range(board_count):
            codes, spi_error = boards[board]
            values = [codes[channel] * 5000.0 / 4096.0 for channel in IV_TO_CHANNEL]
            sheet.append([
                timestamp, sequence, board,
                "读取异常（32字超时）" if spi_error else "正常",
            ] + values)
            if spi_error:
                for cell in sheet[sheet.max_row]:
                    cell.fill = error_fill
                    cell.font = error_font
        workbook.save(self.export_path)

    def _serial_error(self, detail: str) -> None:
        self._disconnect()
        messagebox.showerror("串口错误", detail)

    def _close(self) -> None:
        self._disconnect()
        self.root.destroy()


def main() -> None:
    root = tk.Tk()
    Ads7953App(root)
    root.mainloop()


if __name__ == "__main__":
    main()
