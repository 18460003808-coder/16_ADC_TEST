import unittest
from pathlib import Path
from tempfile import TemporaryDirectory

from openpyxl import Workbook, load_workbook

from ads7953_all_channels_gui import AllChannelsApp
from ads7953_all_channels_gui import all_channel_rows
from ads7953_scope_gui import ScopeApp


class AllChannelMappingTests(unittest.TestCase):
    def test_global_iv_mapping_across_boards(self):
        boards = {
            0: list(range(16)),
            1: [100 + value for value in range(16)],
        }
        rows = list(all_channel_rows(boards, 2))
        self.assertEqual(32, len(rows))
        self.assertEqual((0, 0, 0, 14, 14), rows[0][:5])
        self.assertEqual((15, 0, 15, 0, 0), rows[15][:5])
        self.assertEqual((16, 1, 0, 14, 114), rows[16][:5])
        self.assertEqual((31, 1, 15, 0, 100), rows[31][:5])

    def test_all_channels_excel_uses_iv_rows_and_time_columns(self):
        boards = {
            0: (list(range(16)), False),
            1: ([100 + value for value in range(16)], False),
        }
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "all.xlsx"
            app = AllChannelsApp.__new__(AllChannelsApp)
            app.export_path = path
            app._append_excel(7, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("接口", sheet.cell(1, 1).value)
            self.assertEqual("IV0", sheet.cell(2, 1).value)
            self.assertEqual("IV319", sheet.cell(321, 1).value)
            self.assertIn("Seq 7", sheet.cell(1, 2).value)
            self.assertAlmostEqual(14 * 5000.0 / 4096.0, sheet.cell(2, 2).value)
            self.assertIsNone(sheet.cell(34, 2).value)
            self.assertEqual(2, sheet.max_column)
            self.assertEqual(321, sheet.max_row)

            app._append_excel(8, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertIn("Seq 8", sheet.cell(1, 3).value)
            self.assertEqual(3, sheet.max_column)

    def test_all_channels_excel_highlights_failed_board(self):
        boards = {
            0: (list(range(16)), False),
            1: ([0] * 16, True),
        }
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "all_error.xlsx"
            app = AllChannelsApp.__new__(AllChannelsApp)
            app.export_path = path
            app._append_excel(8, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("IV16", sheet.cell(18, 1).value)
            self.assertEqual("00C00000", sheet.cell(18, 2).font.color.rgb)
            self.assertIsNone(sheet.cell(34, 2).value)

    def test_all_channels_preserves_legacy_sheet(self):
        boards = {0: (list(range(16)), False)}
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "legacy.xlsx"
            workbook = Workbook()
            workbook.active.title = "全通道采集数据"
            workbook.active["A1"] = "旧格式"
            workbook.save(path)
            app = AllChannelsApp.__new__(AllChannelsApp)
            app.export_path = path
            app._append_excel(9, boards, 1)
            workbook = load_workbook(path, data_only=True)
            self.assertIn("全通道采集数据", workbook.sheetnames)
            self.assertIn("纵向通道数据", workbook.sheetnames)
            self.assertEqual("纵向通道数据", workbook.active.title)

    def test_scope_excel_exports_only_selected_global_channels(self):
        boards = {
            0: (list(range(16)), False),
            1: ([100 + value for value in range(16)], True),
        }
        with TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "scope.xlsx"
            app = ScopeApp.__new__(ScopeApp)
            app.export_path = path
            app._selected_channels = lambda: [0, 16]
            app._append_excel(8, boards, 2)
            sheet = load_workbook(path, data_only=True).active
            self.assertEqual("IV0", sheet.cell(2, 3).value)
            self.assertEqual("IV16", sheet.cell(3, 3).value)
            self.assertEqual("读取状态", sheet.cell(1, 7).value)
            self.assertEqual("读取异常（32字超时）", sheet.cell(3, 7).value)
            self.assertEqual(3, sheet.max_row)


if __name__ == "__main__":
    unittest.main()
