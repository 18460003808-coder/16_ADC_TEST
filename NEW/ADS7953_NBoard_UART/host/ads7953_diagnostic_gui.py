"""ADS7953 N-board UART diagnostic view with SPI timeout highlighting."""

from __future__ import annotations

import queue
from datetime import datetime
from pathlib import Path
import tkinter as tk

import serial
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ads7953_Nboard_gui import (
    DATA_FRAME_LENGTH,
    END_FRAME_LENGTH,
    HEADER,
    IV_TO_CHANNEL,
    MAX_BOARD_COUNT,
    PAYLOAD_LENGTH,
    VERSION,
    Ads7953App,
)


class DiagnosticFrameParser:
    """Parse the compatible 40-byte frame and retain its SPI error flag."""

    def __init__(self) -> None:
        self.buffer = bytearray()

    def feed(self, data: bytes) -> list[tuple[str, int, int, list[int], bool]]:
        self.buffer.extend(data)
        frames: list[tuple[str, int, int, list[int], bool]] = []
        while True:
            start = self.buffer.find(HEADER)
            if start < 0:
                self.buffer[:] = self.buffer[-1:] if self.buffer[-1:] == HEADER[:1] else b""
                break
            if start:
                del self.buffer[:start]
            if len(self.buffer) < 5:
                break

            is_end = self.buffer[3] == 0xFF
            frame_length = END_FRAME_LENGTH if is_end else DATA_FRAME_LENGTH
            if len(self.buffer) < frame_length:
                break
            frame = bytes(self.buffer[:frame_length])

            if is_end:
                if (frame[2] != VERSION or not 1 <= frame[4] <= MAX_BOARD_COUNT or
                        self._xor(frame[2:7]) != frame[7]):
                    del self.buffer[0]
                    continue
                sequence = frame[5] | (frame[6] << 8)
                frames.append(("end", sequence, frame[4], [], False))
                del self.buffer[:frame_length]
                continue

            if (frame[2] != VERSION or frame[4] != PAYLOAD_LENGTH or
                    frame[3] >= MAX_BOARD_COUNT or self._xor(frame[2:39]) != frame[39]):
                del self.buffer[0]
                continue

            board_id = frame[3]
            sequence = frame[5] | (frame[6] << 8)
            codes = [
                frame[7 + 2 * channel] | ((frame[8 + 2 * channel] & 0x0F) << 8)
                for channel in range(16)
            ]
            spi_error = bool(frame[8] & 0x80)
            frames.append(("data", sequence, board_id, codes, spi_error))
            del self.buffer[:frame_length]
        return frames

    @staticmethod
    def _xor(data: bytes) -> int:
        result = 0
        for value in data:
            result ^= value
        return result


class DiagnosticApp(Ads7953App):
    def __init__(self, root: tk.Tk) -> None:
        self.current_errors = {board: False for board in range(MAX_BOARD_COUNT)}
        super().__init__(root)
        self.root.title("ADS7953 N板UART读取诊断工具")
        self.export_path = Path.cwd() / "ADS7953_NBoard_Diagnostic_data.xlsx"
        self.pending: dict[int, dict[int, tuple[list[int], bool]]] = {}
        self.table.tag_configure(
            "spi_error", foreground="#C00000", background="#FFE6E6")

    def _reader(self) -> None:
        parser = DiagnosticFrameParser()
        port = self.serial_port
        try:
            while port and port.is_open and not self.stop_event.is_set():
                data = port.read(port.in_waiting or 1)
                for frame in parser.feed(data):
                    self.events.put(("frame", frame))
        except serial.SerialException as exc:
            self.events.put(("error", str(exc)))

    def _poll_events(self) -> None:
        try:
            while True:
                kind, payload = self.events.get_nowait()
                if kind == "frame":
                    frame_kind, sequence, value, codes, spi_error = payload  # type: ignore[misc]
                    if frame_kind == "data":
                        boards = self.pending.setdefault(sequence, {})
                        boards[value] = (codes, spi_error)
                        self.status.set(f"序号 {sequence}：已收到 {len(boards)} 块板")
                    else:
                        board_count = value
                        boards = self.pending.get(sequence, {})
                        missing = [
                            board for board in range(board_count)
                            if board not in boards
                        ]
                        if missing:
                            self.status.set(
                                f"序号 {sequence}不完整，缺少板号："
                                + ",".join(str(board) for board in missing))
                            if not self.continuous_active:
                                self.trigger_button.configure(state=tk.NORMAL)
                            continue
                        self._complete_capture(sequence, boards, board_count)
                        self.pending = {
                            seq: values for seq, values in self.pending.items()
                            if seq > sequence
                        }
                elif kind == "continuous_done":
                    self._stop_continuous()
                    self.status.set("已达到设定采集时间，连续采集已停止")
                else:
                    self._serial_error(str(payload))
        except queue.Empty:
            pass
        self.root.after(50, self._poll_events)

    def _complete_capture(
            self, sequence: int,
            boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        self.detected_board_count = board_count
        for board in range(board_count):
            self.current_codes[board], self.current_errors[board] = boards[board]
        self.board_box["values"] = [str(i) for i in range(board_count)]
        if int(self.board_box.get() or 0) >= board_count:
            self.board_box.current(0)
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self._refresh_table()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)

        failed = [str(board) for board in range(board_count) if self.current_errors[board]]
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        if failed:
            self.status.set(
                f"序号 {sequence}：板号 {','.join(failed)} SPI读取异常（32字超时）{suffix}")
        else:
            mode = "连续采集" if self.continuous_active else "单次采集"
            self.status.set(f"序号 {sequence}：{mode}完成，全部板卡读取正常{suffix}")

    def _refresh_table(self) -> None:
        board = int(self.board_box.get() or 0)
        codes = self.current_codes[board]
        row_tags = ("spi_error",) if self.current_errors[board] else ()
        for iv, channel in enumerate(IV_TO_CHANNEL):
            code = codes[channel]
            millivolts = code * 5000.0 / 4096.0
            self.table.item(
                str(iv),
                values=(f"IV{iv}", channel, code,
                        f"{millivolts:.2f}", f"{millivolts / 1000.0:.4f}"),
                tags=row_tags)

    def _append_excel(
            self, sequence: int,
            boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "采集诊断数据"
            sheet.append(
                ["时间", "序号", "板号", "读取状态"]
                + [f"IV{i}(mV)" for i in range(16)])

        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        error_fill = PatternFill(fill_type="solid", fgColor="FFE6E6")
        error_font = Font(color="C00000")
        for board in range(board_count):
            codes, spi_error = boards[board]
            values = [codes[channel] * 5000.0 / 4096.0 for channel in IV_TO_CHANNEL]
            sheet.append([
                timestamp, sequence, board,
                "读取异常（32字超时）" if spi_error else "正常",
            ] + values)
            if spi_error:
                for cell in sheet[sheet.max_row]:
                    cell.fill = error_fill
                    cell.font = error_font
        workbook.save(self.export_path)


def main() -> None:
    root = tk.Tk()
    DiagnosticApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
