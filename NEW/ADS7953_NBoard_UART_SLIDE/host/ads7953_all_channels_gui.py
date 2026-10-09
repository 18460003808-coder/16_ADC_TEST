"""ADS7953 N-board acquisition tool with all channels on one page."""

from datetime import datetime
from pathlib import Path
import queue
import threading
import time
import tkinter as tk
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from PIL import Image, ImageTk

from ads7953_Nboard_gui import (
    Ads7953App, BAUD, IV_TO_CHANNEL, MAX_BOARD_COUNT, SLIDE_DONE,
    build_slide_command,
)


def all_channel_rows(boards: dict[int, list[int]], board_count: int):
    """Return rows in global IV order: board 0 IV0..15, board 1 IV16..31..."""
    for board in range(board_count):
        codes = boards[board]
        for local_iv, adc_channel in enumerate(IV_TO_CHANNEL):
            global_iv = board * 16 + local_iv
            code = codes[adc_channel]
            millivolts = code * 5000.0 / 4096.0
            yield global_iv, board, local_iv, adc_channel, code, millivolts


def grayscale_pixels_from_excel(
        excel_path: Path, minimum_v: float, maximum_v: float
        ) -> tuple[int, int, bytes]:
    """Load the 320-by-N vertical export and return row-major 8-bit pixels."""
    if minimum_v >= maximum_v:
        raise ValueError("最大电压必须大于最小电压。")
    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    try:
        sheet = (workbook["纵向通道数据"]
                 if "纵向通道数据" in workbook.sheetnames
                 else workbook.active)
        rows = sheet.iter_rows(
            min_row=1, max_row=MAX_BOARD_COUNT * 16 + 1,
            values_only=True)
        header = next(rows, None)
        if not header:
            raise ValueError("Excel中没有可成像的数据。")
        populated_columns = [
            index for index, value in enumerate(header[1:], start=1)
            if value is not None
        ]
        width = max(populated_columns, default=0)
        if width < 1:
            raise ValueError("Excel中没有可成像的采集数据列。")

        pixels = bytearray()
        voltage_span = maximum_v - minimum_v
        for global_iv, row in enumerate(rows):
            expected = f"IV{global_iv}"
            if not row or row[0] != expected:
                raise ValueError(
                    "Excel格式不正确：A列必须依次为IV0～IV319。")
            for data_index, value in enumerate(row[1:width + 1], start=1):
                if value is None:
                    voltage = minimum_v
                else:
                    try:
                        voltage = float(value) / 1000.0
                    except (TypeError, ValueError) as exc:
                        raise ValueError(
                            f"IV{global_iv}第{data_index}组数据不是有效数值。") from exc
                normalized = (voltage - minimum_v) / voltage_span
                pixels.append(round(max(0.0, min(1.0, normalized)) * 255.0))
        if len(pixels) != width * MAX_BOARD_COUNT * 16:
            raise ValueError("Excel数据不足320行。")
        return width, MAX_BOARD_COUNT * 16, bytes(pixels)
    finally:
        workbook.close()


class AllChannelsApp(Ads7953App):
    def __init__(self, root: tk.Tk) -> None:
        self.slide_mode: str | None = None
        self.slide_timeout_token = 0
        self.expected_slide_direction = 0
        self.expected_slide_pulses = 0
        self.slide_started_at = 0.0
        self.link_active = False
        self.link_count = 0
        self.link_completed = 0
        self.link_step_pulses = 0
        self.link_step_mm = 0.0
        self.link_interval_ms = 500
        self.link_return_remaining = 0
        self.link_capture_token = 0
        self.link_capture_started_at = 0.0
        self.link_export_pending = False
        self.grayscale_running = False
        self.grayscale_events: queue.Queue[tuple[str, object]] = queue.Queue()
        self.link_captures: list[
            tuple[str, int, dict[int, tuple[list[int], bool]], int, float]
        ] = []
        super().__init__(root)
        self.export_path = Path.cwd() / "ADS7953_AllChannels_data.xlsx"
        self.root.after(50, self._poll_grayscale_events)

    def _build_ui(self) -> None:
        self.root.title("ADS7953 N板全通道采集工具")
        self.root.geometry("1080x960")

        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        ttk.Label(controls, text="串口：").pack(side=tk.LEFT)
        self.port_box = ttk.Combobox(controls, width=16, state="readonly")
        self.port_box.pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(controls, text="刷新", command=self._refresh_ports).pack(side=tk.LEFT)
        self.connect_button = ttk.Button(controls, text="连接", command=self._toggle_connection)
        self.connect_button.pack(side=tk.LEFT, padx=5)
        self.trigger_button = ttk.Button(
            controls, text="采集", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(
            controls, text="选择普通采集Excel",
            command=self._select_export).pack(side=tk.LEFT, padx=5)

        repeat = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat.pack(fill=tk.X)
        ttk.Label(repeat, text="连续触发间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat, textvariable=self.interval_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(repeat, text="采集时间(s，0=不限)：").pack(side=tk.LEFT)
        self.duration_var = tk.StringVar(value="10")
        ttk.Entry(repeat, textvariable=self.duration_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        self.continuous_button = ttk.Button(
            repeat, text="开始连续采集", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=5)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat, text="连续采集时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)

        slide_debug = ttk.LabelFrame(self.root, text="滑轨单独调试", padding=8)
        slide_debug.pack(fill=tk.X, padx=8, pady=(0, 6))
        ttk.Label(slide_debug, text="方向：").pack(side=tk.LEFT)
        self.slide_direction_var = tk.StringVar(value="正向（dir=1）")
        self.slide_direction_box = ttk.Combobox(
            slide_debug, width=15, state="readonly",
            values=("正向（dir=1）", "反向（dir=0）"),
            textvariable=self.slide_direction_var)
        self.slide_direction_box.pack(side=tk.LEFT, padx=(0, 10))
        ttk.Label(slide_debug, text="距离(mm，0.5～300)：").pack(side=tk.LEFT)
        self.slide_distance_var = tk.StringVar(value="1.0")
        ttk.Entry(slide_debug, textvariable=self.slide_distance_var, width=8).pack(
            side=tk.LEFT, padx=(0, 10))
        self.slide_button = ttk.Button(
            slide_debug, text="滑动", state=tk.DISABLED,
            command=self._start_manual_slide)
        self.slide_button.pack(side=tk.LEFT)
        self.slide_indicator = tk.Label(
            slide_debug, text="滑轨空闲", width=12,
            foreground="white", background="#2E8B57")
        self.slide_indicator.pack(side=tk.LEFT, padx=15)

        linked = ttk.LabelFrame(
            self.root, text="ADC采集与滑轨联动（完成后另存为）", padding=8)
        linked.pack(fill=tk.X, padx=8, pady=(0, 6))
        ttk.Label(linked, text="每步距离(mm)：").pack(side=tk.LEFT)
        self.link_distance_var = tk.StringVar(value="1.0")
        ttk.Entry(linked, textvariable=self.link_distance_var, width=7).pack(
            side=tk.LEFT, padx=(0, 8))
        ttk.Label(linked, text="采集间隔(ms)：").pack(side=tk.LEFT)
        self.link_interval_var = tk.StringVar(value="500")
        ttk.Entry(linked, textvariable=self.link_interval_var, width=7).pack(
            side=tk.LEFT, padx=(0, 8))
        ttk.Label(linked, text="次数N：").pack(side=tk.LEFT)
        self.link_count_var = tk.StringVar(value="10")
        ttk.Entry(linked, textvariable=self.link_count_var, width=7).pack(
            side=tk.LEFT, padx=(0, 8))
        self.link_button = ttk.Button(
            linked, text="开始联动", state=tk.DISABLED,
            command=self._start_linked_run)
        self.link_button.pack(side=tk.LEFT, padx=5)
        self.link_progress_text = tk.StringVar(value="未运行")
        ttk.Label(linked, textvariable=self.link_progress_text).pack(side=tk.LEFT, padx=12)

        grayscale = ttk.LabelFrame(self.root, text="320×N灰度成像", padding=8)
        grayscale.pack(fill=tk.X, padx=8, pady=(0, 6))
        self.image_excel_var = tk.StringVar(value="")
        ttk.Entry(
            grayscale, textvariable=self.image_excel_var,
            state="readonly", width=42).pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(
            grayscale, text="选择Excel",
            command=self._select_image_excel).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(grayscale, text="黑色(V)：").pack(side=tk.LEFT)
        self.image_min_v_var = tk.StringVar(value="0")
        ttk.Entry(grayscale, textvariable=self.image_min_v_var, width=6).pack(
            side=tk.LEFT, padx=(0, 6))
        ttk.Label(grayscale, text="白色(V)：").pack(side=tk.LEFT)
        self.image_max_v_var = tk.StringVar(value="5")
        ttk.Entry(grayscale, textvariable=self.image_max_v_var, width=6).pack(
            side=tk.LEFT, padx=(0, 8))
        self.grayscale_button = ttk.Button(
            grayscale, text="生成灰度图",
            command=self._generate_grayscale_image)
        self.grayscale_button.pack(side=tk.LEFT)

        info = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        info.pack(fill=tk.X)
        self.sequence_text = tk.StringVar(value="采集序号：--")
        self.channel_count_text = tk.StringVar(value="等待识别板卡数量")
        ttk.Label(info, textvariable=self.sequence_text).pack(side=tk.LEFT)
        ttk.Label(info, textvariable=self.channel_count_text).pack(side=tk.LEFT, padx=20)

        self.status = tk.StringVar(value="未连接")
        ttk.Label(self.root, textvariable=self.status, padding=(8, 0, 8, 6)).pack(fill=tk.X)

        table_frame = ttk.Frame(self.root, padding=(8, 0, 8, 8))
        table_frame.pack(fill=tk.BOTH, expand=True)
        columns = ("iv", "board", "local_iv", "channel", "code", "mv", "voltage")
        self.table = ttk.Treeview(table_frame, columns=columns, show="headings")
        headings = ("全局接口", "板号", "板内接口", "ADC通道", "原始码", "毫伏", "电压(V)")
        widths = (100, 70, 100, 100, 100, 130, 130)
        for column, heading, width in zip(columns, headings, widths):
            self.table.heading(column, text=heading)
            self.table.column(column, width=width, anchor=tk.CENTER)
        scrollbar = ttk.Scrollbar(table_frame, orient=tk.VERTICAL, command=self.table.yview)
        self.table.configure(yscrollcommand=scrollbar.set)
        self.table.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.table.tag_configure(
            "spi_error", foreground="#C00000", background="#FFE6E6")

    def _toggle_connection(self) -> None:
        super()._toggle_connection()
        self._update_action_controls()

    def _disconnect(self) -> None:
        self.slide_timeout_token += 1
        self.link_capture_token += 1
        self.slide_mode = None
        self.link_active = False
        super()._disconnect()
        if hasattr(self, "slide_indicator"):
            self._set_slide_indicator(False)
            self.link_progress_text.set(
                "等待保存" if self.link_export_pending else "未运行")
            self._update_action_controls()

    def _toggle_continuous(self) -> None:
        super()._toggle_continuous()
        self._update_action_controls()

    def _stop_continuous(self, update_status: bool = True) -> None:
        super()._stop_continuous(update_status)
        if hasattr(self, "slide_button"):
            self._update_action_controls()

    def _update_action_controls(self) -> None:
        connected = bool(self.serial_port and self.serial_port.is_open)
        exclusive_busy = self.slide_mode is not None or self.link_active
        normal_state = tk.NORMAL if connected and not exclusive_busy else tk.DISABLED
        if not self.continuous_active:
            self.trigger_button.configure(state=normal_state)
        self.continuous_button.configure(state=normal_state)
        slide_state = tk.NORMAL if connected and not exclusive_busy and not self.continuous_active else tk.DISABLED
        self.slide_button.configure(state=slide_state)
        if self.link_export_pending:
            self.link_button.configure(text="保存联动数据", state=tk.NORMAL)
        else:
            self.link_button.configure(text="开始联动", state=slide_state)

    def _set_slide_indicator(self, moving: bool) -> None:
        if moving:
            self.slide_indicator.configure(text="滑轨运行中", background="#C00000")
        else:
            self.slide_indicator.configure(text="滑轨空闲", background="#2E8B57")

    def _select_image_excel(self) -> None:
        selected = filedialog.askopenfilename(
            title="选择320×N采集表格",
            filetypes=[("Excel 工作簿", "*.xlsx")])
        if selected:
            self.image_excel_var.set(selected)
            self.status.set(f"灰度成像数据源：{selected}")

    def _generate_grayscale_image(self) -> None:
        if self.grayscale_running:
            return
        source_text = self.image_excel_var.get().strip()
        if not source_text:
            messagebox.showwarning("未选择Excel", "请先选择已有或刚生成的Excel文件。")
            return
        source_path = Path(source_text)
        if not source_path.is_file():
            messagebox.showwarning("文件不存在", str(source_path))
            return
        try:
            minimum_v = float(self.image_min_v_var.get())
            maximum_v = float(self.image_max_v_var.get())
            if minimum_v >= maximum_v:
                raise ValueError("最大电压必须大于最小电压。")
        except ValueError as exc:
            messagebox.showerror("灰度成像失败", str(exc))
            return

        self.grayscale_running = True
        self.grayscale_button.configure(state=tk.DISABLED)
        self.status.set("正在后台读取Excel并计算320×N灰度数据……")
        threading.Thread(
            target=self._grayscale_worker,
            args=(source_path, minimum_v, maximum_v),
            daemon=True).start()

    def _grayscale_worker(
            self, source_path: Path,
            minimum_v: float, maximum_v: float) -> None:
        try:
            result = grayscale_pixels_from_excel(
                source_path, minimum_v, maximum_v)
            self.grayscale_events.put((
                "ready", (source_path, minimum_v, maximum_v, *result)))
        except Exception as exc:
            self.grayscale_events.put(("error", str(exc)))

    def _poll_grayscale_events(self) -> None:
        try:
            while True:
                kind, payload = self.grayscale_events.get_nowait()
                self.grayscale_running = False
                self.grayscale_button.configure(state=tk.NORMAL)
                if kind == "ready":
                    self._finish_grayscale_image(payload)  # type: ignore[arg-type]
                else:
                    self.status.set("灰度成像失败")
                    messagebox.showerror("灰度成像失败", str(payload))
        except queue.Empty:
            pass
        self.root.after(50, self._poll_grayscale_events)

    def _finish_grayscale_image(
            self, result: tuple[Path, float, float, int, int, bytes]) -> None:
        source_path, minimum_v, maximum_v, width, height, pixels = result
        image = Image.frombytes("L", (width, height), pixels)
        preview_width = min(1000, max(200, width * 4))
        preview_height = 640
        preview = image.resize(
            (preview_width, preview_height), Image.Resampling.NEAREST)
        window = tk.Toplevel(self.root)
        window.title(f"灰度预览：320×{width}")
        photo = ImageTk.PhotoImage(preview)
        label = ttk.Label(window, image=photo)
        label.image = photo
        label.pack(padx=8, pady=8)
        export_status = tk.StringVar(value="当前仅预览，尚未导出PNG")
        ttk.Label(
            window,
            text=(f"纵向IV0～IV319，横向1～{width}次采集；"
                  f"黑={minimum_v:g} V，白={maximum_v:g} V"),
            justify=tk.CENTER).pack(padx=8, pady=(0, 8))
        ttk.Label(
            window, textvariable=export_status,
            justify=tk.CENTER).pack(padx=8, pady=(0, 8))
        actions = ttk.Frame(window)
        actions.pack(pady=(0, 10))
        ttk.Button(
            actions, text="导出PNG",
            command=lambda: self._export_grayscale_png(
                image, source_path, export_status)).pack(side=tk.LEFT, padx=5)
        ttk.Button(
            actions, text="关闭预览",
            command=window.destroy).pack(side=tk.LEFT, padx=5)
        self.status.set("灰度预览已生成；需要保存时点击预览窗口中的“导出PNG”")

    def _export_grayscale_png(
            self, image: Image.Image, source_path: Path,
            export_status: tk.StringVar) -> None:
        output = filedialog.asksaveasfilename(
            title="导出灰度PNG",
            defaultextension=".png",
            filetypes=[("PNG 图像", "*.png")],
            initialdir=str(source_path.parent),
            initialfile=f"{source_path.stem}_grayscale.png")
        if not output:
            export_status.set("已取消导出；预览仍保留")
            return
        try:
            image.save(output)
        except OSError as exc:
            messagebox.showerror("灰度图保存失败", str(exc))
            return
        export_status.set(f"已导出：{output}")
        self.status.set(f"灰度图已导出：{output}")

    @staticmethod
    def _distance_to_pulses(
            value: str, maximum_mm: float = 2.0) -> tuple[float, int]:
        distance_mm = float(value)
        if not 0.5 <= distance_mm <= maximum_mm:
            raise ValueError(f"滑动距离必须在0.5～{maximum_mm:g} mm之间。")
        pulse_count = round(distance_mm * 80.0)
        return pulse_count / 80.0, pulse_count

    def _start_manual_slide(self) -> None:
        try:
            actual_mm, pulse_count = self._distance_to_pulses(
                self.slide_distance_var.get(), maximum_mm=300.0)
        except ValueError as exc:
            messagebox.showwarning("参数错误", str(exc))
            return
        direction = 1 if self.slide_direction_var.get().startswith("正向") else 0
        label = "正向" if direction else "反向"
        if self._send_slide(direction, pulse_count, "manual"):
            self.status.set(
                f"滑轨{label}移动 {actual_mm:g} mm（{pulse_count}个PWM），等待完成……")

    def _send_slide(self, direction: int, pulse_count: int, mode: str) -> bool:
        port = self.serial_port
        if not port or not port.is_open:
            return False
        try:
            port.write(build_slide_command(direction, pulse_count))
        except Exception as exc:
            self._serial_error(str(exc))
            return False
        self.slide_mode = mode
        self.expected_slide_direction = direction
        self.expected_slide_pulses = pulse_count
        self.slide_started_at = time.monotonic()
        self.slide_timeout_token += 1
        token = self.slide_timeout_token
        timeout_ms = int(pulse_count * 1000 / 2000 + 3000)
        self.root.after(timeout_ms, lambda: self._slide_timeout(token))
        self._set_slide_indicator(True)
        self._update_action_controls()
        return True

    def _slide_timeout(self, token: int) -> None:
        if token != self.slide_timeout_token or self.slide_mode is None:
            return
        self.slide_mode = None
        self.link_active = False
        self._set_slide_indicator(False)
        self.link_progress_text.set("已中止")
        self.status.set("滑轨完成应答超时：请检查PWM/方向输出和串口通信")
        self._update_action_controls()

    def _handle_aux_frame(
            self, frame: tuple[str, int, int, list[int], bool]) -> None:
        frame_kind, status_code, direction, values, _unused = frame
        if frame_kind != "slide" or status_code != SLIDE_DONE or not values:
            return
        pulse_count = values[0]
        if self.slide_mode is None:
            self._set_slide_indicator(False)
            return
        if (direction != self.expected_slide_direction or
                pulse_count != self.expected_slide_pulses):
            self.slide_mode = None
            self.link_active = False
            self._set_slide_indicator(False)
            self.link_progress_text.set("已中止")
            self.status.set("收到的滑轨完成帧与本次命令不一致")
            self._update_action_controls()
            return

        mode = self.slide_mode
        self.slide_mode = None
        self.slide_timeout_token += 1
        self._set_slide_indicator(False)

        if mode == "manual":
            label = "正向" if direction else "反向"
            self.status.set(f"滑轨{label}移动完成，共{pulse_count}个PWM")
            self._update_action_controls()
        elif mode == "forward":
            if self.link_completed >= self.link_count:
                self._start_link_return()
            else:
                elapsed_ms = (time.monotonic() - self.link_capture_started_at) * 1000.0
                delay_ms = max(0, int(self.link_interval_ms - elapsed_ms))
                self.status.set(
                    f"联动 {self.link_completed}/{self.link_count}：移动完成，"
                    f"{delay_ms} ms后再次采集")
                self.root.after(delay_ms, self._send_link_capture)
        elif mode == "return":
            self.link_return_remaining -= pulse_count
            if self.link_return_remaining > 0:
                self._send_return_chunk()
            else:
                self._finish_linked_run()

    def _start_linked_run(self) -> None:
        if self.link_export_pending:
            self._prompt_linked_export()
            return
        try:
            actual_mm, step_pulses = self._distance_to_pulses(
                self.link_distance_var.get())
            interval_ms = int(self.link_interval_var.get())
            count = int(self.link_count_var.get())
        except ValueError as exc:
            messagebox.showwarning("参数错误", str(exc))
            return
        if not 500 <= interval_ms <= 5000:
            messagebox.showwarning("参数错误", "采集间隔必须在500～5000 ms之间。")
            return
        if not 1 <= count <= 65535:
            messagebox.showwarning("参数错误", "移动次数N必须在1～65535之间。")
            return

        self.link_active = True
        self.link_count = count
        self.link_completed = 0
        self.link_step_pulses = step_pulses
        self.link_step_mm = actual_mm
        self.link_interval_ms = interval_ms
        self.link_return_remaining = 0
        self.link_export_pending = False
        self.link_captures = []
        self.link_progress_text.set(f"0/{count}")
        self.pending.clear()
        self._update_action_controls()
        self.status.set(
            f"联动开始：每步{actual_mm:g} mm，间隔{interval_ms} ms，共{count}次")
        self._send_link_capture()

    def _send_link_capture(self) -> None:
        if not self.link_active or self.slide_mode is not None:
            return
        port = self.serial_port
        if not port or not port.is_open:
            return
        try:
            port.write(b"T")
        except Exception as exc:
            self._serial_error(str(exc))
            return
        self.pending.clear()
        self.link_capture_started_at = time.monotonic()
        self.link_capture_token += 1
        token = self.link_capture_token
        self.root.after(6000, lambda: self._linked_capture_timeout(token))
        self.status.set(
            f"联动 {self.link_completed + 1}/{self.link_count}：正在采集ADC……")

    def _linked_capture_timeout(self, token: int) -> None:
        if token != self.link_capture_token or not self.link_active:
            return
        self.link_active = False
        self.link_progress_text.set("已中止")
        self.status.set("联动采集超时，流程已中止；请确认当前位置后手动复位")
        self._update_action_controls()

    def _start_link_return(self) -> None:
        self.link_return_remaining = self.link_count * self.link_step_pulses
        self.status.set(
            f"采集完成，滑轨开始反向复位，共{self.link_return_remaining}个PWM")
        self._send_return_chunk()

    def _send_return_chunk(self) -> None:
        chunk = min(self.link_return_remaining, 0xFFFF)
        if not self._send_slide(0, chunk, "return"):
            self.link_active = False
            self.link_progress_text.set("复位失败")
            self._update_action_controls()

    def _finish_linked_run(self) -> None:
        completed = self.link_count
        self.link_active = False
        self.link_export_pending = True
        self.link_progress_text.set(f"完成 {completed}/{completed}，等待保存")
        self.status.set(
            f"联动完成：采集{completed}次，滑轨已复位，请命名并保存新Excel")
        self._update_action_controls()
        self._prompt_linked_export()

    def _prompt_linked_export(self) -> None:
        default_name = datetime.now().strftime("ADS7953_SLIDE_%Y%m%d_%H%M%S.xlsx")
        selected = filedialog.asksaveasfilename(
            title="命名并保存本次联动数据",
            defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")],
            initialfile=default_name)
        if not selected:
            self.link_progress_text.set("等待保存")
            self.status.set(
                "滑轨已复位，本次联动数据仍保留；点击“保存联动数据”可再次命名")
            self._update_action_controls()
            return
        try:
            export_path = Path(selected)
            self._write_linked_excel(export_path)
        except Exception as exc:
            self.link_progress_text.set("导出失败，等待重试")
            self.status.set(f"滑轨已复位，但Excel导出失败：{exc}")
            self._update_action_controls()
            return
        completed = len(self.link_captures)
        self.link_export_pending = False
        self.link_captures = []
        self.image_excel_var.set(str(export_path))
        self.link_button.configure(text="开始联动")
        self.link_progress_text.set(f"完成 {completed}/{completed}")
        self.status.set(
            f"联动完成：采集{completed}次，滑轨已复位，新Excel已保存到 {export_path}")
        self._update_action_controls()

    def _write_linked_excel(self, export_path: Path) -> None:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "纵向通道数据"
        sheet.cell(1, 1, "接口")
        for global_iv in range(MAX_BOARD_COUNT * 16):
            sheet.cell(global_iv + 2, 1, f"IV{global_iv}")
        sheet.freeze_panes = "B2"
        sheet.column_dimensions["A"].width = 12

        error_fill = PatternFill(fill_type="solid", fgColor="FFE6E6")
        error_font = Font(color="C00000")
        for data_column, capture in enumerate(self.link_captures, start=2):
            timestamp, sequence, boards, board_count, position_mm = capture
            sheet.cell(
                1, data_column,
                f"{timestamp} | Seq {sequence} | Pos {position_mm:g} mm")
            sheet.column_dimensions[
                sheet.cell(1, data_column).column_letter].width = 32
            for global_iv in range(MAX_BOARD_COUNT * 16):
                board, local_iv = divmod(global_iv, 16)
                cell = sheet.cell(global_iv + 2, data_column)
                if board >= board_count:
                    continue
                codes, spi_error = boards[board]
                adc_channel = IV_TO_CHANNEL[local_iv]
                cell.value = codes[adc_channel] * 5000.0 / 4096.0
                cell.number_format = "0.00"
                if spi_error:
                    cell.fill = error_fill
                    cell.font = error_font
        workbook.save(export_path)

    def _complete_capture(
            self, sequence: int, boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        if self.link_active:
            self.link_capture_token += 1
            self.detected_board_count = board_count
            for board in range(board_count):
                self.current_codes[board], self.current_errors[board] = boards[board]
            self.current_sequence = sequence
            self.sequence_text.set(f"采集序号：{sequence}")
            self.channel_count_text.set(f"{board_count}块板，共{board_count * 16}路")
            self._refresh_table()

            position_mm = self.link_completed * self.link_step_mm
            snapshot = {
                board: (codes.copy(), spi_error)
                for board, (codes, spi_error) in boards.items()
            }
            timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
            self.link_captures.append(
                (timestamp, sequence, snapshot, board_count, position_mm))
            self.link_completed += 1
            self.link_progress_text.set(
                f"{self.link_completed}/{self.link_count}，位置{position_mm:g} mm")
            self.status.set(
                f"联动 {self.link_completed}/{self.link_count}：采集完成，滑轨正向移动")
            if not self._send_slide(1, self.link_step_pulses, "forward"):
                self.link_active = False
                self.link_progress_text.set("移动失败")
                self._update_action_controls()
            return

        self.detected_board_count = board_count
        for board in range(board_count):
            self.current_codes[board], self.current_errors[board] = boards[board]
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self.channel_count_text.set(f"{board_count}块板，共{board_count * 16}路")
        self._refresh_table()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        failed = [str(board) for board in range(board_count) if self.current_errors[board]]
        if failed:
            self.status.set(
                f"序号 {sequence}：板号 {','.join(failed)} SPI读取异常（32字超时）{suffix}")
        else:
            mode = "连续采集" if self.continuous_active else "单次采集"
            self.status.set(
                f"序号 {sequence}：{mode}完成，{board_count * 16}路，全部正常{suffix}")

    def _refresh_table(self) -> None:
        self.table.delete(*self.table.get_children())
        boards = {
            board: self.current_codes[board]
            for board in range(self.detected_board_count)
        }
        for global_iv, board, local_iv, adc_channel, code, millivolts in all_channel_rows(
                boards, self.detected_board_count):
            tags = ("spi_error",) if self.current_errors[board] else ()
            self.table.insert("", tk.END, values=(
                f"IV{global_iv}", board, f"IV{local_iv}", adc_channel, code,
                f"{millivolts:.2f}", f"{millivolts / 1000.0:.4f}"), tags=tags)

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择全通道数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel输出：{self.export_path}")

    def _append_excel(
            self, sequence: int, boards: dict[int, tuple[list[int], bool]],
            board_count: int, position_mm: float | None = None) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            if "纵向通道数据" in workbook.sheetnames:
                sheet = workbook["纵向通道数据"]
            else:
                # Preserve an existing legacy horizontal-format worksheet.
                sheet = workbook.create_sheet("纵向通道数据")
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "纵向通道数据"
        workbook.active = workbook.index(sheet)

        if sheet.cell(1, 1).value != "接口":
            sheet.cell(1, 1, "接口")
            for global_iv in range(MAX_BOARD_COUNT * 16):
                sheet.cell(global_iv + 2, 1, f"IV{global_iv}")
            sheet.freeze_panes = "B2"
            sheet.column_dimensions["A"].width = 12

        data_column = max(2, sheet.max_column + 1)
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        heading = f"{timestamp} | Seq {sequence}"
        if position_mm is not None:
            heading += f" | Pos {position_mm:g} mm"
        sheet.cell(1, data_column, heading)
        sheet.column_dimensions[sheet.cell(1, data_column).column_letter].width = 24

        error_fill = PatternFill(fill_type="solid", fgColor="FFE6E6")
        error_font = Font(color="C00000")
        for global_iv in range(MAX_BOARD_COUNT * 16):
            board, local_iv = divmod(global_iv, 16)
            cell = sheet.cell(global_iv + 2, data_column)
            if board >= board_count:
                cell.value = None
                continue
            codes, spi_error = boards[board]
            adc_channel = IV_TO_CHANNEL[local_iv]
            cell.value = codes[adc_channel] * 5000.0 / 4096.0
            cell.number_format = "0.00"
            if spi_error:
                cell.fill = error_fill
                cell.font = error_font
        workbook.save(self.export_path)
        if hasattr(self, "image_excel_var"):
            self.image_excel_var.set(str(self.export_path))


def main() -> None:
    root = tk.Tk()
    AllChannelsApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
