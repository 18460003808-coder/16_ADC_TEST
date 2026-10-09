"""ADS7953 N-board live waveform viewer."""

from collections import deque
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, ttk

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from ads7953_Nboard_gui import Ads7953App, IV_TO_CHANNEL, MAX_BOARD_COUNT


HISTORY_POINTS = 200
TRACE_COLORS = (
    "#00d4ff", "#ffcc00", "#ff5f57", "#7ee787", "#d2a8ff", "#ffa657",
    "#79c0ff", "#f778ba", "#a5d6ff", "#56d364", "#e3b341", "#ff7b72",
)


class ScopeApp(Ads7953App):
    def __init__(self, root: tk.Tk) -> None:
        self.sample_numbers: deque[int] = deque(maxlen=HISTORY_POINTS)
        self.histories = [deque(maxlen=HISTORY_POINTS) for _ in range(MAX_BOARD_COUNT * 16)]
        self.error_histories = [
            deque(maxlen=HISTORY_POINTS) for _ in range(MAX_BOARD_COUNT * 16)]
        self.sample_number = 0
        super().__init__(root)
        self.export_path = Path.cwd() / "ADS7953_Scope_data.xlsx"

    def _build_ui(self) -> None:
        self.root.title("ADS7953 N板实时波形工具")
        self.root.geometry("1180x760")

        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        ttk.Label(controls, text="串口：").pack(side=tk.LEFT)
        self.port_box = ttk.Combobox(controls, width=16, state="readonly")
        self.port_box.pack(side=tk.LEFT, padx=(0, 5))
        ttk.Button(controls, text="刷新", command=self._refresh_ports).pack(side=tk.LEFT)
        self.connect_button = ttk.Button(controls, text="连接", command=self._toggle_connection)
        self.connect_button.pack(side=tk.LEFT, padx=5)
        self.trigger_button = ttk.Button(
            controls, text="单次采集", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="选择Excel", command=self._select_export).pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="清除波形", command=self._clear_waveforms).pack(side=tk.LEFT, padx=5)

        repeat = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat.pack(fill=tk.X)
        ttk.Label(repeat, text="采样间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat, textvariable=self.interval_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(repeat, text="观察时间(s，0=不限)：").pack(side=tk.LEFT)
        self.duration_var = tk.StringVar(value="10")
        ttk.Entry(repeat, textvariable=self.duration_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        self.continuous_button = ttk.Button(
            repeat, text="开始实时波形", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=5)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat, text="实时观察时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)

        self.sequence_text = tk.StringVar(value="采集序号：--")
        self.status = tk.StringVar(value="未连接")
        info = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        info.pack(fill=tk.X)
        ttk.Label(info, textvariable=self.sequence_text).pack(side=tk.LEFT)
        ttk.Label(info, textvariable=self.status).pack(side=tk.LEFT, padx=20)

        content = ttk.Frame(self.root, padding=(8, 0, 8, 8))
        content.pack(fill=tk.BOTH, expand=True)
        selector = ttk.LabelFrame(content, text="接口通道（可多选）", padding=6)
        selector.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 8))
        self.channel_list = tk.Listbox(
            selector, width=16, selectmode=tk.EXTENDED, exportselection=False)
        channel_scroll = ttk.Scrollbar(selector, orient=tk.VERTICAL, command=self.channel_list.yview)
        self.channel_list.configure(yscrollcommand=channel_scroll.set)
        self.channel_list.pack(side=tk.LEFT, fill=tk.Y)
        channel_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self.channel_list.insert(tk.END, "IV0")
        self.channel_list.selection_set(0)
        self.channel_list.bind("<<ListboxSelect>>", lambda _event: self._draw_scope())

        self.canvas = tk.Canvas(content, background="#10151c", highlightthickness=0)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.canvas.bind("<Configure>", lambda _event: self._draw_scope())

    def _complete_capture(
            self, sequence: int, boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        count_changed = self.detected_board_count != board_count
        self.detected_board_count = board_count
        self.sample_number += 1
        self.sample_numbers.append(self.sample_number)
        for board in range(board_count):
            codes, spi_error = boards[board]
            self.current_codes[board] = codes
            self.current_errors[board] = spi_error
            for local_iv, adc_channel in enumerate(IV_TO_CHANNEL):
                global_iv = board * 16 + local_iv
                self.histories[global_iv].append(
                    codes[adc_channel] * 5.0 / 4096.0)
                self.error_histories[global_iv].append(spi_error)
        if count_changed:
            self._update_channel_list(board_count * 16)
        self._update_channel_error_colors(board_count * 16)
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self._draw_scope()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        failed = [str(board) for board in range(board_count) if self.current_errors[board]]
        if failed:
            self.status.set(
                f"板号 {','.join(failed)} SPI读取异常（红色采样点，32字超时）{suffix}")
        else:
            mode = "实时观察" if self.continuous_active else "单次采集"
            self.status.set(f"{mode}：{board_count}块板，{board_count * 16}路，全部正常{suffix}")

    def _update_channel_list(self, channel_count: int) -> None:
        selected = set(self.channel_list.curselection())
        self.channel_list.delete(0, tk.END)
        for global_iv in range(channel_count):
            self.channel_list.insert(tk.END, f"IV{global_iv}")
        valid = sorted(index for index in selected if index < channel_count)
        if not valid and channel_count:
            valid = [0]
        for index in valid:
            self.channel_list.selection_set(index)

    def _update_channel_error_colors(self, channel_count: int) -> None:
        for global_iv in range(channel_count):
            board = global_iv // 16
            color = "#C00000" if self.current_errors[board] else "#000000"
            self.channel_list.itemconfig(global_iv, foreground=color)

    def _selected_channels(self) -> list[int]:
        return list(self.channel_list.curselection())

    def _stop_continuous(self, update_status: bool = True) -> None:
        super()._stop_continuous(update_status)
        self.continuous_button.configure(text="开始实时波形")

    def _draw_scope(self) -> None:
        canvas = self.canvas
        canvas.delete("all")
        width = max(canvas.winfo_width(), 320)
        height = max(canvas.winfo_height(), 240)
        left, top, right, bottom = 58, 42, width - 18, height - 40
        plot_width = max(right - left, 1)
        plot_height = max(bottom - top, 1)

        for volt in range(6):
            y = bottom - volt * plot_height / 5.0
            canvas.create_line(left, y, right, y, fill="#2d3742")
            canvas.create_text(left - 10, y, text=f"{volt}V", fill="#aab6c3", anchor=tk.E)
        canvas.create_line(left, top, left, bottom, fill="#d0d7de")
        canvas.create_line(left, bottom, right, bottom, fill="#d0d7de")
        canvas.create_text((left + right) / 2, height - 14, text="最近200次采集", fill="#aab6c3")

        selected = self._selected_channels()
        for trace_index, global_iv in enumerate(selected):
            values = list(self.histories[global_iv])
            errors = list(self.error_histories[global_iv])
            if not values:
                continue
            color = TRACE_COLORS[trace_index % len(TRACE_COLORS)]
            if len(values) == 1:
                x = right
                y = bottom - max(0.0, min(5.0, values[0])) * plot_height / 5.0
                canvas.create_oval(x - 2, y - 2, x + 2, y + 2, fill=color, outline=color)
            else:
                points = []
                start_index = HISTORY_POINTS - len(values)
                for index, voltage in enumerate(values):
                    x = left + (start_index + index) * plot_width / (HISTORY_POINTS - 1)
                    y = bottom - max(0.0, min(5.0, voltage)) * plot_height / 5.0
                    points.extend((x, y))
                canvas.create_line(*points, fill=color, width=2)
            start_index = HISTORY_POINTS - len(values)
            for index, (voltage, spi_error) in enumerate(zip(values, errors)):
                if not spi_error:
                    continue
                x = left + (start_index + index) * plot_width / (HISTORY_POINTS - 1)
                y = bottom - max(0.0, min(5.0, voltage)) * plot_height / 5.0
                canvas.create_oval(
                    x - 4, y - 4, x + 4, y + 4,
                    fill="#FF3030", outline="#FFFFFF", width=1)
            legend_x = left + (trace_index % 6) * 145
            legend_y = 14 + (trace_index // 6) * 18
            canvas.create_text(
                legend_x, legend_y, anchor=tk.W, fill=color,
                text=f"IV{global_iv}  {values[-1]:.4f}V")

    def _clear_waveforms(self) -> None:
        self.sample_numbers.clear()
        for history in self.histories:
            history.clear()
        for history in self.error_histories:
            history.clear()
        self._draw_scope()
        self.status.set("波形记录已清除")

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择波形数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel输出：{self.export_path}")

    def _append_excel(
            self, sequence: int, boards: dict[int, tuple[list[int], bool]],
            board_count: int) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "波形通道数据"
            sheet.append([
                "时间", "序号", "接口", "板号", "板内接口",
                "ADC通道", "读取状态", "原始码", "毫伏"])
        if sheet.cell(1, 7).value != "读取状态":
            sheet.insert_cols(7)
            sheet.cell(1, 7, "读取状态")
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        error_fill = PatternFill(fill_type="solid", fgColor="FFE6E6")
        error_font = Font(color="C00000")
        for global_iv in self._selected_channels():
            board, local_iv = divmod(global_iv, 16)
            if board >= board_count:
                continue
            adc_channel = IV_TO_CHANNEL[local_iv]
            codes, spi_error = boards[board]
            code = codes[adc_channel]
            millivolts = code * 5000.0 / 4096.0
            sheet.append([
                timestamp, sequence, f"IV{global_iv}", board, f"IV{local_iv}",
                adc_channel, "读取异常（32字超时）" if spi_error else "正常",
                code, millivolts])
            if spi_error:
                for cell in sheet[sheet.max_row]:
                    cell.fill = error_fill
                    cell.font = error_font
        workbook.save(self.export_path)


def main() -> None:
    root = tk.Tk()
    ScopeApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
