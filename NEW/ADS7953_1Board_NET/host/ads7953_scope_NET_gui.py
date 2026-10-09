"""ADS7953 N-board UDP live waveform viewer."""

from collections import deque
from datetime import datetime
from pathlib import Path
import tkinter as tk
from tkinter import filedialog, ttk

from openpyxl import Workbook, load_workbook

from ads7953_Nboard_NET_gui import Ads7953App, IV_TO_CHANNEL, MAX_BOARD_COUNT


DEFAULT_HISTORY_POINTS = 200
MAX_HISTORY_POINTS = 5000
TRACE_COLORS = (
    "#00d4ff", "#ffcc00", "#ff5f57", "#7ee787", "#d2a8ff", "#ffa657",
    "#79c0ff", "#f778ba", "#a5d6ff", "#56d364", "#e3b341", "#ff7b72",
)


def validate_scale_values(
        x_points_text: str, x_tick_text: str,
        y_min_text: str, y_max_text: str,
        y_tick_text: str) -> tuple[int, int, float, float, float]:
    """Parse and validate user-entered scope scales."""
    x_points = int(x_points_text)
    x_tick = int(x_tick_text)
    y_min = float(y_min_text)
    y_max = float(y_max_text)
    y_tick = float(y_tick_text)
    if not 10 <= x_points <= MAX_HISTORY_POINTS:
        raise ValueError(f"横轴显示点数必须为10～{MAX_HISTORY_POINTS}")
    if not 1 <= x_tick <= x_points:
        raise ValueError("横轴刻度间隔必须为1～显示点数")
    if y_max <= y_min:
        raise ValueError("纵轴最大值必须大于最小值")
    if y_tick <= 0 or (y_max - y_min) / y_tick > 100:
        raise ValueError("纵轴刻度间隔必须大于0，且刻度数量不能超过100")
    return x_points, x_tick, y_min, y_max, y_tick


class ScopeApp(Ads7953App):
    def __init__(self, root: tk.Tk) -> None:
        self.sample_numbers: deque[int] = deque(maxlen=MAX_HISTORY_POINTS)
        self.histories = [
            deque(maxlen=MAX_HISTORY_POINTS) for _ in range(MAX_BOARD_COUNT * 16)]
        self.sample_number = 0
        self.x_points = DEFAULT_HISTORY_POINTS
        self.x_tick = 20
        self.y_min = 0.0
        self.y_max = 5.0
        self.y_tick = 1.0
        self.channel_vars: dict[int, tk.BooleanVar] = {}
        super().__init__(root)
        self.export_path = Path.cwd() / "ADS7953_Scope_NET_data.xlsx"

    def _build_ui(self) -> None:
        self.root.title("ADS7953 N板网口实时波形工具")
        self.root.geometry("1180x760")

        controls = ttk.Frame(self.root, padding=8)
        controls.pack(fill=tk.X)
        self._build_network_controls(controls)
        self.trigger_button = ttk.Button(
            controls, text="单次采集", state=tk.DISABLED, command=self._trigger)
        self.trigger_button.pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="选择Excel", command=self._select_export).pack(side=tk.LEFT, padx=5)
        ttk.Button(controls, text="清除波形", command=self._clear_waveforms).pack(side=tk.LEFT, padx=5)

        repeat = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        repeat.pack(fill=tk.X)
        ttk.Label(repeat, text="采样间隔(ms)：").pack(side=tk.LEFT)
        self.interval_var = tk.StringVar(value="500")
        ttk.Entry(repeat, textvariable=self.interval_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        ttk.Label(repeat, text="观察时间(s，0=不限)：").pack(side=tk.LEFT)
        self.duration_var = tk.StringVar(value="10")
        ttk.Entry(repeat, textvariable=self.duration_var, width=8).pack(side=tk.LEFT, padx=(0, 8))
        self.continuous_button = ttk.Button(
            repeat, text="开始实时波形", state=tk.DISABLED,
            command=self._toggle_continuous)
        self.continuous_button.pack(side=tk.LEFT, padx=5)
        self.continuous_excel_var = tk.BooleanVar(value=False)
        ttk.Checkbutton(
            repeat, text="实时观察时写Excel",
            variable=self.continuous_excel_var).pack(side=tk.LEFT, padx=8)

        scale = ttk.LabelFrame(self.root, text="显示刻度", padding=(8, 4))
        scale.pack(fill=tk.X, padx=8, pady=(0, 6))
        self.x_points_var = tk.StringVar(value=str(self.x_points))
        self.x_tick_var = tk.StringVar(value=str(self.x_tick))
        self.y_min_var = tk.StringVar(value=f"{self.y_min:g}")
        self.y_max_var = tk.StringVar(value=f"{self.y_max:g}")
        self.y_tick_var = tk.StringVar(value=f"{self.y_tick:g}")
        scale_items = (
            ("横轴显示点数：", self.x_points_var, 7),
            ("横轴刻度间隔：", self.x_tick_var, 7),
            ("纵轴最小值(V)：", self.y_min_var, 7),
            ("纵轴最大值(V)：", self.y_max_var, 7),
            ("纵轴刻度(V)：", self.y_tick_var, 7),
        )
        for label, variable, entry_width in scale_items:
            ttk.Label(scale, text=label).pack(side=tk.LEFT)
            ttk.Entry(scale, textvariable=variable, width=entry_width).pack(
                side=tk.LEFT, padx=(0, 8))
        ttk.Button(scale, text="应用刻度", command=self._apply_scales).pack(
            side=tk.LEFT, padx=4)

        self.sequence_text = tk.StringVar(value="采集序号：--")
        self.status = tk.StringVar(value="未连接")
        info = ttk.Frame(self.root, padding=(8, 0, 8, 6))
        info.pack(fill=tk.X)
        ttk.Label(info, textvariable=self.sequence_text).pack(side=tk.LEFT)
        ttk.Label(info, textvariable=self.status).pack(side=tk.LEFT, padx=20)

        content = ttk.Frame(self.root, padding=(8, 0, 8, 8))
        content.pack(fill=tk.BOTH, expand=True)
        selector = ttk.LabelFrame(content, text="接口通道（勾选显示）", padding=6)
        selector.pack(side=tk.LEFT, fill=tk.Y, padx=(0, 8))
        self.channel_canvas = tk.Canvas(
            selector, width=125, highlightthickness=0, borderwidth=0)
        channel_scroll = ttk.Scrollbar(
            selector, orient=tk.VERTICAL, command=self.channel_canvas.yview)
        self.channel_canvas.configure(yscrollcommand=channel_scroll.set)
        self.channel_frame = ttk.Frame(self.channel_canvas)
        self.channel_window = self.channel_canvas.create_window(
            (0, 0), window=self.channel_frame, anchor=tk.NW)
        self.channel_frame.bind(
            "<Configure>",
            lambda _event: self.channel_canvas.configure(
                scrollregion=self.channel_canvas.bbox("all")))
        self.channel_canvas.bind(
            "<Configure>",
            lambda event: self.channel_canvas.itemconfigure(
                self.channel_window, width=event.width))
        self.channel_canvas.pack(side=tk.LEFT, fill=tk.Y)
        channel_scroll.pack(side=tk.RIGHT, fill=tk.Y)
        self._update_channel_list(1)

        self.canvas = tk.Canvas(content, background="#10151c", highlightthickness=0)
        self.canvas.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        self.canvas.bind("<Configure>", lambda _event: self._draw_scope())

    def _complete_capture(
            self, sequence: int, boards: dict[int, list[int]],
            board_count: int) -> None:
        count_changed = self.detected_board_count != board_count
        self.detected_board_count = board_count
        self.sample_number += 1
        self.sample_numbers.append(self.sample_number)
        for board in range(board_count):
            self.current_codes[board] = boards[board]
            for local_iv, adc_channel in enumerate(IV_TO_CHANNEL):
                global_iv = board * 16 + local_iv
                self.histories[global_iv].append(
                    boards[board][adc_channel] * 5.0 / 4096.0)
        if count_changed:
            self._update_channel_list(board_count * 16)
        self.current_sequence = sequence
        self.sequence_text.set(f"采集序号：{sequence}")
        self._draw_scope()
        wrote_excel = not self.continuous_active or self.continuous_excel_var.get()
        if wrote_excel:
            self._append_excel(sequence, boards, board_count)
        if not self.continuous_active:
            self.trigger_button.configure(state=tk.NORMAL)
        suffix = f"，已写入 {self.export_path}" if wrote_excel else ""
        mode = "实时观察" if self.continuous_active else "单次采集"
        self.status.set(f"{mode}：{board_count}块板，{board_count * 16}路{suffix}")

    def _update_channel_list(self, channel_count: int) -> None:
        selected = set(self._selected_channels())
        if not selected and channel_count:
            selected = {0}
        for child in self.channel_frame.winfo_children():
            child.destroy()
        self.channel_vars.clear()
        for global_iv in range(channel_count):
            variable = tk.BooleanVar(value=global_iv in selected)
            self.channel_vars[global_iv] = variable
            ttk.Checkbutton(
                self.channel_frame, text=f"IV{global_iv}", variable=variable,
                command=self._draw_scope).pack(anchor=tk.W, fill=tk.X)

    def _selected_channels(self) -> list[int]:
        return [index for index, variable in self.channel_vars.items()
                if variable.get()]

    def _apply_scales(self) -> None:
        try:
            values = validate_scale_values(
                self.x_points_var.get(), self.x_tick_var.get(),
                self.y_min_var.get(), self.y_max_var.get(),
                self.y_tick_var.get())
        except (TypeError, ValueError) as exc:
            self.status.set(f"刻度设置错误：{exc}")
            return
        self.x_points, self.x_tick, self.y_min, self.y_max, self.y_tick = values
        self._draw_scope()
        self.status.set("显示刻度已更新")

    def _stop_continuous(self, update_status: bool = True) -> None:
        super()._stop_continuous(update_status)
        self.continuous_button.configure(text="开始实时波形")

    def _draw_scope(self) -> None:
        canvas = self.canvas
        canvas.delete("all")
        width = max(canvas.winfo_width(), 320)
        height = max(canvas.winfo_height(), 240)
        left, top, right, bottom = 58, 42, width - 18, height - 40
        plot_width = max(right - left, 1)
        plot_height = max(bottom - top, 1)

        y_value = self.y_min
        while y_value <= self.y_max + self.y_tick * 1e-9:
            y = bottom - (y_value - self.y_min) * plot_height / (
                self.y_max - self.y_min)
            canvas.create_line(left, y, right, y, fill="#2d3742")
            canvas.create_text(
                left - 10, y, text=f"{y_value:g}V",
                fill="#aab6c3", anchor=tk.E)
            y_value += self.y_tick
        canvas.create_line(left, top, left, bottom, fill="#d0d7de")
        canvas.create_line(left, bottom, right, bottom, fill="#d0d7de")
        window_start = max(1, self.sample_number - self.x_points + 1)
        denominator = max(self.x_points - 1, 1)
        for offset in range(0, self.x_points, self.x_tick):
            x = left + offset * plot_width / denominator
            canvas.create_line(x, top, x, bottom, fill="#24303b")
            canvas.create_text(
                x, bottom + 14, text=str(window_start + offset),
                fill="#aab6c3", anchor=tk.N)
        canvas.create_text(
            (left + right) / 2, height - 8,
            text=f"采集序号（最近{self.x_points}次）", fill="#aab6c3")

        selected = self._selected_channels()
        for trace_index, global_iv in enumerate(selected):
            values = list(self.histories[global_iv])[-self.x_points:]
            if not values:
                continue
            color = TRACE_COLORS[trace_index % len(TRACE_COLORS)]
            if len(values) == 1:
                x = right
                clamped = max(self.y_min, min(self.y_max, values[0]))
                y = bottom - (clamped - self.y_min) * plot_height / (
                    self.y_max - self.y_min)
                canvas.create_oval(x - 2, y - 2, x + 2, y + 2, fill=color, outline=color)
            else:
                points = []
                start_index = self.x_points - len(values)
                for index, voltage in enumerate(values):
                    x = left + (start_index + index) * plot_width / denominator
                    clamped = max(self.y_min, min(self.y_max, voltage))
                    y = bottom - (clamped - self.y_min) * plot_height / (
                        self.y_max - self.y_min)
                    points.extend((x, y))
                canvas.create_line(*points, fill=color, width=2)
            legend_x = left + (trace_index % 6) * 145
            legend_y = 14 + (trace_index // 6) * 18
            canvas.create_text(
                legend_x, legend_y, anchor=tk.W, fill=color,
                text=f"IV{global_iv}  {values[-1]:.4f}V")

    def _clear_waveforms(self) -> None:
        self.sample_numbers.clear()
        for history in self.histories:
            history.clear()
        self._draw_scope()
        self.status.set("波形记录已清除")

    def _select_export(self) -> None:
        selected = filedialog.asksaveasfilename(
            title="选择波形数据文件", defaultextension=".xlsx",
            filetypes=[("Excel 工作簿", "*.xlsx")], initialfile=self.export_path.name)
        if selected:
            self.export_path = Path(selected)
            self.status.set(f"Excel输出：{self.export_path}")

    def _append_excel(
            self, sequence: int, boards: dict[int, list[int]],
            board_count: int) -> None:
        if self.export_path.exists():
            workbook = load_workbook(self.export_path)
            sheet = workbook.active
        else:
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "波形通道数据"
            sheet.append([
                "时间", "序号", "接口", "板号", "板内接口",
                "ADC通道", "原始码", "毫伏"])
        timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S.%f")[:-3]
        for global_iv in self._selected_channels():
            board, local_iv = divmod(global_iv, 16)
            if board >= board_count:
                continue
            adc_channel = IV_TO_CHANNEL[local_iv]
            code = boards[board][adc_channel]
            millivolts = code * 5000.0 / 4096.0
            sheet.append([
                timestamp, sequence, f"IV{global_iv}", board, f"IV{local_iv}",
                adc_channel, code, millivolts])
        workbook.save(self.export_path)


def main() -> None:
    root = tk.Tk()
    ScopeApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
